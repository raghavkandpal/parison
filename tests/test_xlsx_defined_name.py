import json
import tempfile
import unittest
import zipfile
from copy import deepcopy
from datetime import date
from pathlib import Path
from unittest.mock import patch

from jsonschema import Draft202012Validator

from parison.core import ParisonError, compare, draft_recipe, load_schema, run_suite, validate_inputs, verify_bundle

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.workbook.defined_name import DefinedName
except ImportError:
    Workbook = None
    DefinedName = None
    load_workbook = None

RECIPE = {
    "recipe_version": 1,
    "comparison_mode": "keyed",
    "keys": ["id"],
    "scope": {
        "snapshot": "xlsx-defined-name-test",
        "cutoff": "2026-10-19T00:00:00Z",
        "filters": [],
        "completeness": "full",
        "expected_empty": False,
    },
    "identity": {"null_keys": "reject", "duplicates": "reject"},
    "nulls_equal": True,
    "input_projection": "semantic-scalar-v1",
    "columns": {
        "id": {"type": "string", "comparison": "exact"},
        "value": {"type": "integer", "comparison": "exact"},
        "day": {"type": "date", "comparison": "exact"},
    },
    "output": {"sensitivity": "summary"},
}


@unittest.skipUnless(Workbook, "openpyxl optional dependency is not installed")
class XlsxDefinedName(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.recipe = self.root / "recipe.json"
        self.workbook = self.root / "records.xlsx"
        self.recipe.write_text(json.dumps(RECIPE), encoding="utf-8")
        self._write_workbook("'Sales Data'!$A$2:$C$4")

    def tearDown(self):
        self.tmp.cleanup()

    @property
    def source(self):
        return f"xlsx:{self.workbook}#defined-name=OrdersExport"

    def _write_workbook(self, definition, *, local_sheet_id=None, name="OrdersExport"):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Sales Data"
        for row in (
            ["ignored"],
            ["id", "value", "day"],
            ["001", 10, date(2026, 10, 17)],
            ["002", 20, date(2026, 10, 18)],
        ):
            worksheet.append(row)
        workbook.defined_names.add(DefinedName(name, localSheetId=local_sheet_id, attr_text=definition))
        workbook.save(self.workbook)
        workbook.close()

    def _rewrite_archive(self, transform):
        replacement = self.root / "replacement.xlsx"
        with zipfile.ZipFile(self.workbook) as source, zipfile.ZipFile(replacement, "w") as target:
            for member in source.infolist():
                target.writestr(member, transform(member.filename, source.read(member)))
        replacement.replace(self.workbook)

    def test_exact_defined_name_all_modes_engines_metadata_preflight_and_suite(self):
        result = compare(self.recipe, self.source, self.source)
        self.assertEqual(result["outcome"], "PASS")
        self.assertNotIn("defined_name", result["inputs"]["baseline"])
        self.assertNotIn("worksheet", result["inputs"]["baseline"])
        self.assertNotIn("range", result["inputs"]["baseline"])
        self.assertEqual(draft_recipe(self.source, self.source)["input_projection"], "semantic-scalar-v1")

        raw = deepcopy(RECIPE)
        raw["output"] = {"sensitivity": "raw"}
        self.recipe.write_text(json.dumps(raw), encoding="utf-8")
        metadata = compare(self.recipe, self.source, self.source)["inputs"]["baseline"]
        self.assertEqual(metadata["defined_name"], "OrdersExport")
        self.assertEqual(metadata["worksheet"], "Sales Data")
        self.assertEqual(metadata["range"], "A2:C4")

        for version, schema_name in ((1, "preflight-v19"), (2, "preflight-v20"), (3, "preflight-v21")):
            recipe = deepcopy(RECIPE)
            if version == 2:
                recipe.update(
                    recipe_version=2,
                    comparison_mode="aggregate",
                    group_by=["id"],
                    measures={"total": {"operator": "sum", "column": "value", "nulls": "reject"}},
                )
                for key in ("keys", "identity", "nulls_equal"):
                    recipe.pop(key)
            elif version == 3:
                recipe.update(recipe_version=3, comparison_mode="multiset")
                for key in ("keys", "identity", "nulls_equal"):
                    recipe.pop(key)
            self.recipe.write_text(json.dumps(recipe), encoding="utf-8")
            with self.subTest(version=version):
                for engine in ("memory", "sqlite"):
                    self.assertEqual(compare(self.recipe, self.source, self.source, engine=engine)["outcome"], "PASS")
                preflight = validate_inputs(self.recipe, self.source, self.source, validate_records=True)
                self.assertEqual(preflight["schema_version"], 18 + version)
                Draft202012Validator(load_schema(schema_name)).validate(preflight)

        self.recipe.write_text(json.dumps(RECIPE), encoding="utf-8")
        plan = self.root / "suite.json"
        plan.write_text(
            json.dumps(
                {
                    "suite_version": 2,
                    "cases": [
                        {
                            "id": "defined-name",
                            "recipe": "recipe.json",
                            "baseline": "xlsx:records.xlsx#defined-name=OrdersExport",
                            "candidate": "xlsx:records.xlsx#defined-name=OrdersExport",
                            "execution": {"engine": "sqlite", "max_temp_bytes": 1_000_000},
                        }
                    ],
                }
            ),
            encoding="utf-8",
        )
        output = self.root / "suite-output"
        self.assertEqual(run_suite(plan, output)["outcome"], "PASS")
        self.assertEqual(verify_bundle(output)["kind"], "suite")

    def test_field_projection_recipe_can_use_defined_name_locator(self):
        for version, schema_version in ((4, 22), (5, 23), (6, 24)):
            recipe = deepcopy(RECIPE)
            recipe["recipe_version"] = version
            recipe["field_projection"] = {
                "version": 1,
                "baseline": {"kind": "direct-v1", "fields": {name: name for name in recipe["columns"]}},
                "candidate": {"kind": "direct-v1", "fields": {name: name for name in recipe["columns"]}},
            }
            if version == 5:
                recipe.update(
                    comparison_mode="aggregate",
                    group_by=["id"],
                    measures={"total": {"operator": "sum", "column": "value", "nulls": "reject"}},
                )
                for key in ("keys", "identity", "nulls_equal"):
                    recipe.pop(key)
            elif version == 6:
                recipe.update(comparison_mode="multiset")
                for key in ("keys", "identity", "nulls_equal"):
                    recipe.pop(key)
            self.recipe.write_text(json.dumps(recipe), encoding="utf-8")
            with self.subTest(version=version):
                for engine in ("memory", "sqlite"):
                    self.assertEqual(compare(self.recipe, self.source, self.source, engine=engine)["outcome"], "PASS")
                preflight = validate_inputs(self.recipe, self.source, self.source, validate_records=True)
                self.assertEqual(preflight["schema_version"], schema_version)
                validator = Draft202012Validator(load_schema(f"preflight-v{schema_version}"))
                validator.validate(preflight)
                wrong_projection = deepcopy(preflight)
                wrong_projection["field_projection"]["baseline"]["kind"] = "json-pointer-v1"
                self.assertFalse(validator.is_valid(wrong_projection))

    def test_rejects_scope_case_reserved_dynamic_relative_external_and_union(self):
        with self.assertRaisesRegex(ParisonError, "does not exist"):
            compare(self.recipe, self.source.replace("OrdersExport", "ordersexport"), self.source)
        with self.assertRaisesRegex(ParisonError, "non-reserved"):
            compare(self.recipe, self.source.replace("OrdersExport", "_xlnm.Print_Area"), self.source)

        cases = (
            ("'Sales Data'!$A$2:$C$4", 0, "workbook-scoped"),
            ("42", None, "static workbook range"),
            ("OFFSET('Sales Data'!$A$2,0,0,3,3)", None, "static workbook range"),
            ("'Sales Data'!A2:C4", None, "absolute bounded rectangle"),
            ("'[1]Sales Data'!$A$2:$C$4", None, "local worksheet"),
            ("'Sales Data'!$A$2:$C$4,'Sales Data'!$A$2:$C$4", None, "one destination"),
            ("'Sales Data'!$A:$C", None, "absolute bounded rectangle"),
        )
        for definition, local_sheet_id, message in cases:
            self._write_workbook(definition, local_sheet_id=local_sheet_id)
            with self.subTest(definition=definition), self.assertRaisesRegex(ParisonError, message):
                compare(self.recipe, self.source, self.source)

    def test_rejects_duplicate_and_conflicting_raw_metadata(self):
        def duplicate(_, data):
            if _ == "xl/workbook.xml":
                marker = b"</definedNames>"
                declaration = b"<definedName name=\"OrdersExport\">'Sales Data'!$A$2:$C$4</definedName>"
                return data.replace(marker, declaration + marker)
            return data

        self._rewrite_archive(duplicate)
        with self.assertRaisesRegex(ParisonError, "unique within scope"):
            compare(self.recipe, self.source, self.source)

        self._write_workbook("'Sales Data'!$A$2:$C$4")

        def conflict(name, data):
            if name == "xl/workbook.xml":
                marker = b"</definedNames>"
                declaration = (
                    b'<definedName name="OrdersExport" localSheetId="0">\'Sales Data\'!$A$2:$C$4</definedName>'
                )
                return data.replace(marker, declaration + marker)
            return data

        self._rewrite_archive(conflict)
        with self.assertRaisesRegex(ParisonError, "global and local names conflict"):
            compare(self.recipe, self.source, self.source)

        self._write_workbook("'Sales Data'!$A$2:$C$4")

        def invalid_scope(name, data):
            if name == "xl/workbook.xml":
                marker = b"</definedNames>"
                declaration = b'<definedName name="LocalOnly" localSheetId="-1">$A$1</definedName>'
                return data.replace(marker, declaration + marker)
            return data

        self._rewrite_archive(invalid_scope)
        with self.assertRaisesRegex(ParisonError, "local sheet scope is invalid"):
            compare(self.recipe, self.source, self.source)

        self._write_workbook("'Sales Data'!$A$2:$C$4")

        def equivalent_scopes(name, data):
            if name == "xl/workbook.xml":
                marker = b"</definedNames>"
                declarations = (
                    b'<definedName name="LocalOnly" localSheetId="0">$A$1</definedName>'
                    b'<definedName name="LocalOnly" localSheetId="00">$A$1</definedName>'
                )
                return data.replace(marker, declarations + marker)
            return data

        self._rewrite_archive(equivalent_scopes)
        with self.assertRaisesRegex(ParisonError, "unique within scope"):
            compare(self.recipe, self.source, self.source)

    def test_reuses_formula_error_zip_and_row_guards(self):
        workbook = load_workbook(self.workbook)
        workbook["Sales Data"]["B3"] = "=1+1"
        workbook.save(self.workbook)
        workbook.close()
        with self.assertRaisesRegex(ParisonError, "formula cell"):
            compare(self.recipe, self.source, self.source)

        self._write_workbook("'Sales Data'!$A$2:$C$4")
        workbook = load_workbook(self.workbook)
        workbook["Sales Data"]["B3"] = "#DIV/0!"
        workbook.save(self.workbook)
        workbook.close()
        with self.assertRaisesRegex(ParisonError, "error cell"):
            compare(self.recipe, self.source, self.source)

        self._write_workbook("'Sales Data'!$A$2:$C$4")
        with self.assertRaisesRegex(ParisonError, "row count.*exceeds limit 1"):
            validate_inputs(self.recipe, self.source, self.source, validate_records=True, max_rows=1)
        with self.assertRaisesRegex(ParisonError, "decoded input size"):
            compare(self.recipe, self.source, self.source, max_decoded_bytes=1)

    def test_defined_name_count_is_bounded_before_openpyxl_load(self):
        def excessive_names(name, data):
            if name == "xl/workbook.xml":
                marker = b"</definedNames>"
                declarations = b"".join(
                    f'<definedName name="N{index}">$A$1</definedName>'.encode() for index in range(10_001)
                )
                return data.replace(marker, declarations + marker)
            return data

        self._rewrite_archive(excessive_names)
        with patch("openpyxl.load_workbook", side_effect=AssertionError("openpyxl loaded before raw audit")):
            with self.assertRaisesRegex(ParisonError, "exceeds 10000 declarations"):
                compare(self.recipe, self.source, self.source)

    def test_defined_name_audit_rejects_dtd_entities_before_openpyxl_load(self):
        def entity_definition(name, data):
            if name == "xl/workbook.xml":
                declaration = b'<!DOCTYPE workbook [<!ENTITY hostile "OrdersExport">]>'
                return data.replace(b"<workbook ", declaration + b"<workbook ").replace(
                    b'name="OrdersExport"', b'name="&hostile;"'
                )
            return data

        self._rewrite_archive(entity_definition)
        with patch("openpyxl.load_workbook", side_effect=AssertionError("openpyxl loaded before raw audit")):
            with self.assertRaisesRegex(ParisonError, "invalid defined-name metadata"):
                compare(self.recipe, self.source, self.source)

    def test_defined_name_audit_rejects_utf16_entity_metadata_before_openpyxl_load(self):
        def utf16_entity_definition(name, data):
            if name == "xl/workbook.xml":
                text = data.decode("utf-8").replace('encoding="UTF-8"', 'encoding="UTF-16"')
                text = text.replace(
                    "<workbook ", '<!DOCTYPE workbook [<!ENTITY hostile "OrdersExport">]><workbook '
                ).replace('name="OrdersExport"', 'name="&hostile;"')
                return text.encode("utf-16")
            return data

        self._rewrite_archive(utf16_entity_definition)
        with patch("openpyxl.load_workbook", side_effect=AssertionError("openpyxl loaded before raw audit")):
            with self.assertRaisesRegex(ParisonError, "ASCII-compatible XML encoding"):
                compare(self.recipe, self.source, self.source)


if __name__ == "__main__":
    unittest.main()
