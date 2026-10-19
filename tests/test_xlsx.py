import csv
import json
import tempfile
import unittest
import warnings
import zipfile
from copy import deepcopy
from datetime import date, datetime
from pathlib import Path

from jsonschema import Draft202012Validator

from parison.core import (
    ParisonError,
    compare,
    draft_recipe,
    load_schema,
    load_suite,
    publish,
    validate_inputs,
    verify_bundle,
)

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.worksheet.table import Table, TableFormula
except ImportError:
    Workbook = None
    Table = None
    TableFormula = None
    load_workbook = None

RECIPE = {
    "recipe_version": 1,
    "comparison_mode": "keyed",
    "keys": ["id"],
    "scope": {
        "snapshot": "xlsx-test",
        "cutoff": "2026-10-18T00:00:00Z",
        "filters": [],
        "completeness": "full",
        "expected_empty": False,
    },
    "identity": {"null_keys": "reject", "duplicates": "reject"},
    "nulls_equal": True,
    "input_projection": "semantic-scalar-v1",
    "columns": {
        "id": {"type": "string", "comparison": "exact"},
        "value": {"type": "integer", "comparison": "exact"},
        "day": {"type": "date", "comparison": "exact"},
    },
    "output": {"sensitivity": "summary"},
}


@unittest.skipUnless(Workbook, "openpyxl optional dependency is not installed")
class XlsxInput(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.recipe = self.root / "recipe.json"
        self.recipe.write_text(json.dumps(RECIPE), encoding="utf-8")
        self.workbook = self.root / "records.xlsx"
        self._write_workbook(
            [
                ["ignored"],
                ["id", "value", "day"],
                ["001", 10, date(2026, 10, 17)],
                ["002", 20, date(2026, 10, 18)],
            ]
        )

    def tearDown(self):
        self.tmp.cleanup()

    def _write_workbook(self, rows, table=None, table_ref=None):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Export"
        for row in rows:
            worksheet.append(row)
        if table is not None:
            worksheet.add_table(Table(displayName=table, ref=table_ref))
        workbook.save(self.workbook)
        workbook.close()

    @property
    def source(self):
        return f"xlsx:{self.workbook}#Export!A2:C4"

    @property
    def table_source(self):
        return f"xlsx:{self.workbook}#table=Orders"

    def _write_named_table(self):
        self._write_workbook(
            [
                ["ignored"],
                ["id", "value", "day"],
                ["001", 10, date(2026, 10, 17)],
                ["002", 20, date(2026, 10, 18)],
            ],
            table="Orders",
            table_ref="A2:C4",
        )

    def _rewrite_workbook_archive(self, transform):
        replacement = self.root / "replacement.xlsx"
        with zipfile.ZipFile(self.workbook) as source, zipfile.ZipFile(replacement, "w") as target:
            for member in source.infolist():
                data = transform(member.filename, source.read(member))
                if data is not None:
                    target.writestr(member, data)
        replacement.replace(self.workbook)

    def test_named_table_compares_drafts_and_records_resolved_metadata(self):
        self._write_named_table()
        result = compare(self.recipe, self.table_source, self.table_source)
        self.assertEqual(result["outcome"], "PASS")
        self.assertEqual(compare(self.recipe, self.table_source, self.table_source, engine="sqlite")["outcome"], "PASS")
        self.assertEqual(result["counts"]["matched_exact"], 2)
        metadata = result["inputs"]["baseline"]
        self.assertNotIn("table", metadata)
        self.assertNotIn("worksheet", metadata)
        self.assertNotIn("range", metadata)
        self.assertNotIn(str(self.workbook), json.dumps(result))
        self.assertEqual(draft_recipe(self.table_source, self.table_source)["input_projection"], "semantic-scalar-v1")

        raw_recipe = deepcopy(RECIPE)
        raw_recipe["output"] = {"sensitivity": "raw"}
        self.recipe.write_text(json.dumps(raw_recipe), encoding="utf-8")
        raw_metadata = compare(self.recipe, self.table_source, self.table_source)["inputs"]["baseline"]
        self.assertEqual(raw_metadata["table"], "Orders")
        self.assertEqual(raw_metadata["worksheet"], "Export")
        self.assertEqual(raw_metadata["range"], "A2:C4")
        self.recipe.write_text(json.dumps(RECIPE), encoding="utf-8")

        preflight = validate_inputs(self.recipe, self.table_source, self.table_source, validate_records=True)
        self.assertEqual(preflight["schema_version"], 13)
        Draft202012Validator(load_schema("preflight-v13")).validate(preflight)
        without_table = deepcopy(preflight)
        without_table["inputs"]["baseline"].pop("table")
        without_table["inputs"]["candidate"].pop("table")
        self.assertFalse(Draft202012Validator(load_schema("preflight-v13")).is_valid(without_table))

        open_recipe = dict(RECIPE)
        open_recipe.pop("input_projection")
        self.recipe.write_text(json.dumps(open_recipe), encoding="utf-8")
        with self.assertRaisesRegex(ParisonError, "require input_projection"):
            validate_inputs(self.recipe, self.table_source, self.table_source)

    def test_named_table_normalizes_its_resolved_range(self):
        self._write_named_table()
        raw_recipe = deepcopy(RECIPE)
        raw_recipe["output"] = {"sensitivity": "raw"}
        self.recipe.write_text(json.dumps(raw_recipe), encoding="utf-8")
        workbook = load_workbook(self.workbook)
        workbook["Export"].tables["Orders"].ref = "$A$2:$C$4"
        workbook.save(self.workbook)
        workbook.close()
        result = compare(self.recipe, self.table_source, self.table_source)
        self.assertEqual(result["inputs"]["baseline"]["range"], "A2:C4")

    def test_named_table_supports_all_modes_and_portable_suites(self):
        self._write_named_table()
        aggregate = deepcopy(RECIPE)
        aggregate.update(
            recipe_version=2,
            comparison_mode="aggregate",
            group_by=["id"],
            measures={"value_sum": {"operator": "sum", "column": "value", "nulls": "reject"}},
        )
        for key in ("keys", "identity", "nulls_equal"):
            aggregate.pop(key)
        self.recipe.write_text(json.dumps(aggregate), encoding="utf-8")
        self.assertEqual(compare(self.recipe, self.table_source, self.table_source)["outcome"], "PASS")
        self.assertEqual(compare(self.recipe, self.table_source, self.table_source, engine="sqlite")["outcome"], "PASS")
        aggregate_preflight = validate_inputs(self.recipe, self.table_source, self.table_source)
        self.assertEqual(aggregate_preflight["schema_version"], 14)
        Draft202012Validator(load_schema("preflight-v14")).validate(aggregate_preflight)

        multiset = deepcopy(RECIPE)
        multiset.update(recipe_version=3, comparison_mode="multiset")
        for key in ("keys", "identity", "nulls_equal"):
            multiset.pop(key)
        self.recipe.write_text(json.dumps(multiset), encoding="utf-8")
        self.assertEqual(compare(self.recipe, self.table_source, self.table_source)["outcome"], "PASS")
        self.assertEqual(compare(self.recipe, self.table_source, self.table_source, engine="sqlite")["outcome"], "PASS")
        multiset_preflight = validate_inputs(self.recipe, self.table_source, self.table_source)
        self.assertEqual(multiset_preflight["schema_version"], 15)
        Draft202012Validator(load_schema("preflight-v15")).validate(multiset_preflight)

        suite = self.root / "suite.json"
        portable = "xlsx:records.xlsx#table=Orders"
        suite.write_text(
            json.dumps(
                {
                    "suite_version": 2,
                    "cases": [
                        {"id": "xlsx-table", "recipe": "recipe.json", "baseline": portable, "candidate": portable}
                    ],
                }
            ),
            encoding="utf-8",
        )
        self.assertEqual(load_suite(suite)["cases"][0]["baseline"], self.table_source)

    def test_named_table_rejects_missing_case_variant_and_totals(self):
        self._write_named_table()
        with self.assertRaisesRegex(ParisonError, "table does not exist"):
            compare(self.recipe, self.table_source.replace("Orders", "orders"), self.table_source)

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Export"
        worksheet.append(["id", "value", "day"])
        worksheet.append(["001", 10, date(2026, 10, 17)])
        worksheet.add_table(Table(displayName="Orders", ref="A1:C2", totalsRowShown=True))
        workbook.save(self.workbook)
        workbook.close()
        with self.assertRaisesRegex(ParisonError, "no totals row"):
            compare(self.recipe, self.table_source, self.table_source)

    def test_named_table_rejects_metadata_formula_header_mismatch_and_header_only_ref(self):
        self._write_named_table()
        workbook = load_workbook(self.workbook)
        table = workbook["Export"].tables["Orders"]
        table.tableColumns[1].calculatedColumnFormula = TableFormula(attr_text="SUM([value])")
        workbook.save(self.workbook)
        workbook.close()
        with self.assertRaisesRegex(ParisonError, "formula declaration"):
            compare(self.recipe, self.table_source, self.table_source)

        self._write_named_table()
        workbook = load_workbook(self.workbook)
        workbook["Export"].tables["Orders"].tableColumns[0].name = "wrong"
        workbook.save(self.workbook)
        workbook.close()
        with self.assertRaisesRegex(ParisonError, "do not match stored header"):
            compare(self.recipe, self.table_source, self.table_source)

        self._write_workbook([["id", "value", "day"]], table="Orders", table_ref="A1:C1")
        with self.assertRaisesRegex(ParisonError, "header and data"):
            compare(self.recipe, self.table_source, self.table_source)

    def test_named_table_rejects_broken_relationships_targets_and_duplicate_identities(self):
        workbook = Workbook()
        for index, (name, table_name) in enumerate((("Export", "Orders"), ("Broken", "Other"))):
            worksheet = workbook.active if index == 0 else workbook.create_sheet()
            worksheet.title = name
            worksheet.append(["id", "value", "day"])
            worksheet.append(["001", 10, date(2026, 10, 17)])
            worksheet.add_table(Table(displayName=table_name, ref="A1:C2"))
        workbook.save(self.workbook)
        workbook.close()
        self._rewrite_workbook_archive(
            lambda name, data: None if name == "xl/worksheets/_rels/sheet2.xml.rels" else data
        )
        with self.assertRaisesRegex(ParisonError, "invalid table metadata"):
            compare(self.recipe, self.table_source, self.table_source)

        self._write_named_table()
        self._rewrite_workbook_archive(
            lambda name, data: (
                data.replace(b"/xl/tables/table1.xml", b"../../../evil.xml")
                if name == "xl/worksheets/_rels/sheet1.xml.rels"
                else data
            )
        )
        with self.assertRaisesRegex(ParisonError, "invalid package target"):
            compare(self.recipe, self.table_source, self.table_source)

        workbook = Workbook()
        for index, table_name in enumerate(("Orders", "Other", "Third")):
            worksheet = workbook.active if index == 0 else workbook.create_sheet()
            worksheet.title = f"Sheet{index + 1}"
            worksheet.append(["id", "value", "day"])
            worksheet.append(["001", 10, date(2026, 10, 17)])
            worksheet.add_table(Table(displayName=table_name, ref="A1:C2"))
        workbook.save(self.workbook)
        workbook.close()
        self._rewrite_workbook_archive(
            lambda name, data: data.replace(b"Third", b"Other") if name == "xl/tables/table3.xml" else data
        )
        with self.assertRaisesRegex(ParisonError, "names and identifiers.*unique"):
            compare(self.recipe, self.table_source, self.table_source)

        self._write_named_table()
        self._rewrite_workbook_archive(
            lambda name, data: (
                data.replace(b'<tableParts count="1">', b'<tableParts count="2">')
                if name == "xl/worksheets/sheet1.xml"
                else data
            )
        )
        with self.assertRaisesRegex(ParisonError, "count does not match"):
            compare(self.recipe, self.table_source, self.table_source)

    def test_explicit_range_compares_and_records_bounded_source_metadata(self):
        result = compare(self.recipe, self.source, self.source)
        self.assertEqual(result["outcome"], "PASS")
        self.assertEqual(result["counts"]["matched_exact"], 2)
        metadata = result["inputs"]["baseline"]
        self.assertEqual(metadata["format"], "xlsx")
        self.assertEqual(metadata["worksheet"], "Export")
        self.assertEqual(metadata["range"], "A2:C4")
        self.assertEqual(metadata["compression"], "zip")
        self.assertGreater(metadata["decoded_bytes"], metadata["bytes"])
        self.assertNotIn(str(self.workbook), json.dumps(result))

    def test_xlsx_result_publishes_as_a_verifiable_existing_bundle(self):
        output = self.root / "run"
        result = compare(self.recipe, self.source, self.source)
        publish(output, result, RECIPE)
        verify_bundle(output)
        published = json.loads((output / "result.json").read_text(encoding="utf-8"))
        self.assertEqual(published["outcome"], "PASS")
        self.assertEqual(published["inputs"]["baseline"]["format"], "xlsx")

    def test_draft_recipe_preserves_closed_projection(self):
        draft = draft_recipe(self.source, self.source)
        self.assertEqual(draft["input_projection"], "semantic-scalar-v1")
        self.assertEqual(draft["keys"], ["id"])
        self.assertEqual(draft["columns"]["value"]["type"], "integer")
        self.assertEqual(draft["columns"]["day"]["type"], "date")

    def test_preflight_uses_new_schema_without_changing_released_schema(self):
        result = validate_inputs(self.recipe, self.source, self.source, validate_records=True)
        self.assertEqual(result["schema_version"], 4)
        self.assertEqual(result["inputs"]["baseline"]["worksheet"], "Export")
        Draft202012Validator(load_schema("preflight-v4")).validate(result)

    def test_all_modes_and_mixed_text_use_the_same_values(self):
        baseline = self.root / "records.csv"
        with baseline.open("w", newline="", encoding="utf-8") as handle:
            csv.writer(handle).writerows(
                (["id", "value", "day"], ["001", "10", "2026-10-17"], ["002", "20", "2026-10-18"])
            )
        self.assertEqual(compare(self.recipe, baseline, self.source)["outcome"], "PASS")

        aggregate = deepcopy(RECIPE)
        aggregate.update(
            recipe_version=2,
            comparison_mode="aggregate",
            group_by=["id"],
            measures={"value_sum": {"operator": "sum", "column": "value", "nulls": "reject"}},
        )
        for key in ("keys", "identity", "nulls_equal"):
            aggregate.pop(key)
        self.recipe.write_text(json.dumps(aggregate), encoding="utf-8")
        self.assertEqual(compare(self.recipe, baseline, self.source)["outcome"], "PASS")
        aggregate_preflight = validate_inputs(self.recipe, baseline, self.source, validate_records=True)
        Draft202012Validator(load_schema("preflight-v5")).validate(aggregate_preflight)
        aggregate_draft = draft_recipe(self.source, self.source, aggregate=True)
        self.assertEqual(aggregate_draft["comparison_mode"], "aggregate")
        self.assertEqual(aggregate_draft["input_projection"], "semantic-scalar-v1")

        multiset = deepcopy(RECIPE)
        multiset.update(recipe_version=3, comparison_mode="multiset")
        for key in ("keys", "identity", "nulls_equal"):
            multiset.pop(key)
        self.recipe.write_text(json.dumps(multiset), encoding="utf-8")
        self.assertEqual(compare(self.recipe, baseline, self.source)["outcome"], "PASS")
        multiset_preflight = validate_inputs(self.recipe, baseline, self.source, validate_records=True)
        Draft202012Validator(load_schema("preflight-v6")).validate(multiset_preflight)
        multiset_draft = draft_recipe(self.source, self.source, multiset=True)
        self.assertEqual(multiset_draft["comparison_mode"], "multiset")
        self.assertEqual(multiset_draft["input_projection"], "semantic-scalar-v1")

    def test_sqlite_engines_match_memory_results(self):
        memory = compare(self.recipe, self.source, self.source)
        sqlite = compare(self.recipe, self.source, self.source, engine="sqlite")
        self.assertEqual(sqlite["outcome"], memory["outcome"])
        self.assertEqual(sqlite["counts"], memory["counts"])

        multiset = deepcopy(RECIPE)
        multiset.update(recipe_version=3, comparison_mode="multiset")
        for key in ("keys", "identity", "nulls_equal"):
            multiset.pop(key)
        self.recipe.write_text(json.dumps(multiset), encoding="utf-8")
        memory = compare(self.recipe, self.source, self.source)
        sqlite = compare(self.recipe, self.source, self.source, engine="sqlite")
        self.assertEqual(sqlite["outcome"], memory["outcome"])
        self.assertEqual(sqlite["counts"], memory["counts"])

    def test_suite_resolves_portable_xlsx_locators(self):
        suite = self.root / "suite.json"
        portable = "xlsx:records.xlsx#Export!A2:C4"
        suite.write_text(
            json.dumps(
                {
                    "suite_version": 2,
                    "cases": [
                        {
                            "id": "xlsx-values",
                            "recipe": "recipe.json",
                            "baseline": portable,
                            "candidate": portable,
                        }
                    ],
                }
            ),
            encoding="utf-8",
        )
        case = load_suite(suite)["cases"][0]
        self.assertEqual(case["baseline"], f"xlsx:{self.workbook}#Export!A2:C4")
        self.assertEqual(case["candidate"], case["baseline"])

    def test_rejects_implicit_selection_formulas_and_open_projection(self):
        with self.assertRaisesRegex(ParisonError, "XLSX locator"):
            compare(self.recipe, self.workbook, self.source)
        with self.assertRaisesRegex(ParisonError, "worksheet does not exist"):
            compare(self.recipe, self.source.replace("#Export", "#Missing"), self.source)
        with self.assertRaisesRegex(ParisonError, "bounded rectangle"):
            compare(self.recipe, self.source.replace("A2:C4", "A:C"), self.source)

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Export"
        worksheet.append(["id", "value", "day"])
        worksheet.append(["001", "=5+5", date(2026, 10, 17)])
        workbook.save(self.workbook)
        workbook.close()
        formula_source = f"xlsx:{self.workbook}#Export!A1:C2"
        with self.assertRaisesRegex(ParisonError, "formula cell"):
            compare(self.recipe, formula_source, formula_source)

        open_recipe = dict(RECIPE)
        open_recipe.pop("input_projection")
        self.recipe.write_text(json.dumps(open_recipe), encoding="utf-8")
        with self.assertRaisesRegex(ParisonError, "require input_projection"):
            compare(self.recipe, formula_source, formula_source)

    def test_rejects_formula_and_error_cells_in_headers_and_records(self):
        cases = (('="id"', "formula cell"), ("#DIV/0!", "error cell"))
        for value, message in cases:
            for row in (1, 2):
                with self.subTest(value=value, row=row):
                    workbook = Workbook()
                    worksheet = workbook.active
                    worksheet.title = "Export"
                    worksheet.append(["id", "value", "day"])
                    worksheet.append(["001", 10, date(2026, 10, 17)])
                    worksheet.cell(row=row, column=1, value=value)
                    if message == "error cell":
                        worksheet.cell(row=row, column=1).data_type = "e"
                    workbook.save(self.workbook)
                    workbook.close()
                    source = f"xlsx:{self.workbook}#Export!A1:C2"
                    with self.assertRaisesRegex(ParisonError, message):
                        compare(self.recipe, source, source)

    def test_rejects_malformed_duplicate_headers_and_naive_datetimes(self):
        self.workbook.write_bytes(b"not a workbook")
        with self.assertRaisesRegex(ParisonError, "cannot inspect XLSX archive"):
            compare(self.recipe, self.source, self.source)

        self._write_workbook([["id", "id", "day"], ["001", 10, date(2026, 10, 17)]])
        source = f"xlsx:{self.workbook}#Export!A1:C2"
        with self.assertRaisesRegex(ParisonError, "duplicate column names"):
            compare(self.recipe, source, source)

        self._write_workbook([["id", "observed_at"], ["001", datetime(2026, 10, 17, 12)]])
        timestamp_recipe = deepcopy(RECIPE)
        timestamp_recipe["columns"] = {
            "id": {"type": "string", "comparison": "exact"},
            "observed_at": {"type": "timestamp", "comparison": "exact", "timezone": "require-aware"},
        }
        self.recipe.write_text(json.dumps(timestamp_recipe), encoding="utf-8")
        source = f"xlsx:{self.workbook}#Export!A1:B2"
        with self.assertRaisesRegex(ParisonError, "timezone required"):
            compare(self.recipe, source, source)

    def test_rejects_symlink_and_duplicate_archive_members(self):
        link = self.root / "linked.xlsx"
        link.symlink_to(self.workbook)
        linked_source = f"xlsx:{link}#Export!A2:C4"
        with self.assertRaisesRegex(ParisonError, "non-symlink"):
            compare(self.recipe, linked_source, linked_source)

        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            with zipfile.ZipFile(self.workbook, "w") as archive:
                archive.writestr("duplicate", b"one")
                archive.writestr("duplicate", b"two")
        with self.assertRaisesRegex(ParisonError, "duplicate member names"):
            compare(self.recipe, self.source, self.source)

    def test_rejects_row_and_decoded_size_overruns(self):
        with self.assertRaisesRegex(ParisonError, "combined input size"):
            compare(self.recipe, self.source, self.source, max_input_bytes=1)
        with self.assertRaisesRegex(ParisonError, "row count"):
            compare(self.recipe, self.source, self.source, max_rows=1)
        with self.assertRaisesRegex(ParisonError, "decoded input size"):
            compare(self.recipe, self.source, self.source, max_decoded_bytes=1)


if __name__ == "__main__":
    unittest.main()
