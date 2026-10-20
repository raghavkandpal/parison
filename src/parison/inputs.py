"""Local input readers, source identity, validation, and recipe drafting."""

from __future__ import annotations

import csv
import gzip
import hashlib
import io
import json
import re
import sqlite3
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from urllib.parse import parse_qsl

from .core import (
    _FILE_FORMATS,
    ParisonError,
    _checked_policy_sha256,
    _normalize,
    _parse_input,
    _parse_row,
    _source_name,
    load_recipe,
)


def _file_format(path: Path) -> str | None:
    suffix = path.suffix.lower()
    if suffix == ".gz":
        suffix = Path(path.stem).suffix.lower()
        if suffix not in {".csv", ".tsv", ".jsonl", ".ndjson"}:
            return None
    return _FILE_FORMATS.get(suffix)


class _LimitedText:
    def __init__(self, handle, limit: int, path: Path):
        self.handle, self.remaining, self.path = handle, limit, path

    def __enter__(self):
        return self

    def __exit__(self, *args):
        self.handle.close()

    def __iter__(self):
        return self

    def __next__(self):
        line = self.handle.readline(self.remaining + 1)
        if not line:
            raise StopIteration
        size = len(line.encode("utf-8"))
        if size > self.remaining:
            raise ParisonError(f"decoded input changed or exceeds checked size: {self.path}")
        self.remaining -= size
        return line


def _open_text(path: Path, decoded_sizes: dict[Path, int], newline: str | None = None):
    if path.suffix.lower() != ".gz":
        return path.open("r", encoding="utf-8", newline=newline)
    return _LimitedText(gzip.open(path, "rt", encoding="utf-8", newline=newline), decoded_sizes[path], path)


def _validate_headers(path: Path, headers: list[str], recipe: dict[str, Any], side: str) -> None:
    if len(headers) != len(set(headers)):
        raise ParisonError(f"duplicate column names in {path}")
    details = _schema_details(headers, recipe, side)
    missing, extra = details["missing"], details["unexpected"]
    if missing or extra:
        parts = []
        if missing:
            parts.append("missing=" + ",".join(missing))
        if extra:
            parts.append("unexpected=" + ",".join(extra))
        raise ParisonError(f"schema mismatch in {path}: {'; '.join(parts)}")


def _schema_details(headers: list[str], recipe: dict[str, Any], side: str) -> dict[str, Any]:
    physical = {name: _source_name(recipe, name, side) for name in recipe["columns"]}
    header_names = set(headers)
    expected = set(physical.values())
    excluded = set(recipe.get("excluded_columns", {}))
    return {
        "missing": sorted(expected - header_names),
        "unexpected": sorted(header_names - expected - excluded),
        "mapped": {name: source for name, source in physical.items() if source != name},
        "excluded_present": sorted(header_names & excluded),
    }


def _sqlite_source(source: str | Path) -> tuple[Path, str] | None:
    text = str(source)
    if not text.startswith("sqlite:"):
        return None
    locator = text.removeprefix("sqlite:")
    if "?" in locator or "#" not in locator:
        raise ParisonError("SQLite locator must be sqlite:path#table without parameters")
    path_text, table = locator.rsplit("#", 1)
    if not path_text or not table:
        raise ParisonError("SQLite locator requires a database path and table")
    return Path(path_text), table


def _xlsx_source(source: str | Path) -> tuple[Path, str, str | None, str] | None:
    text = str(source)
    if not text.startswith("xlsx:"):
        return None
    locator = text.removeprefix("xlsx:")
    if "?" in locator or "#" not in locator:
        raise ParisonError(
            "XLSX locator must select a worksheet range, exact table name, or exact defined name without parameters"
        )
    path_text, selection = locator.rsplit("#", 1)
    if selection.startswith("table="):
        kind, worksheet, cell_range = "table", None, selection.removeprefix("table=")
        if "!" in cell_range:
            raise ParisonError("XLSX table locator must contain one exact table name")
    elif selection.startswith("defined-name="):
        kind, worksheet, cell_range = "defined-name", None, selection.removeprefix("defined-name=")
        try:
            valid_name = bool(cell_range) and len(cell_range.encode("utf-8")) <= 255
        except UnicodeEncodeError:
            valid_name = False
        if not valid_name or cell_range.casefold().startswith("_xlnm."):
            raise ParisonError("XLSX defined-name locator requires one non-reserved exact name up to 255 bytes")
    elif "!" in selection:
        kind = "range"
        worksheet, cell_range = selection.rsplit("!", 1)
    else:
        raise ParisonError("XLSX locator must select a worksheet range, exact table name, or exact defined name")
    if not path_text or not cell_range or (worksheet is not None and not worksheet):
        raise ParisonError("XLSX locator requires a workbook path and explicit selection")
    if Path(path_text).suffix.lower() != ".xlsx":
        raise ParisonError("XLSX locator requires a .xlsx workbook")
    return Path(path_text), kind, worksheet, cell_range


def _xlsx_table(workbook, table_name: str):
    from openpyxl.packaging.relationship import get_dependents, get_rels_path
    from openpyxl.worksheet.table import Table, TableList
    from openpyxl.xml.constants import REL_NS, SHEET_MAIN_NS
    from openpyxl.xml.functions import fromstring, iterparse

    matches = []
    names, identifiers_seen, targets = set(), set(), set()
    for worksheet in workbook.worksheets:
        tables = TableList()
        try:
            with workbook._archive.open(worksheet._worksheet_path) as source:
                identifiers, declared_counts = [], []
                for _, element in iterparse(source, events=("end",)):
                    if element.tag == f"{{{SHEET_MAIN_NS}}}tablePart":
                        identifiers.append(element.attrib.get(f"{{{REL_NS}}}id"))
                    elif element.tag == f"{{{SHEET_MAIN_NS}}}tableParts":
                        declared_counts.append(element.attrib.get("count"))
                    element.clear()
            if declared_counts:
                if len(declared_counts) != 1 or declared_counts[0] != str(len(identifiers)):
                    raise ValueError("table part count does not match declared metadata")
            elif identifiers:
                raise ValueError("table parts are missing their container")
            if not identifiers:
                continue
            relationships = get_dependents(workbook._archive, get_rels_path(worksheet._worksheet_path))
            relationship_ids = [relationship.Id for relationship in relationships]
            if len(relationship_ids) != len(set(relationship_ids)) or len(identifiers) != len(set(identifiers)):
                raise ValueError("table relationships must be unique")
            for identifier in identifiers:
                relationship = relationships.get(identifier)
                if relationship.Type != Table._rel_type:
                    raise ValueError("table part has an invalid relationship type")
                if not relationship.target.startswith("xl/tables/") or ".." in Path(relationship.target).parts:
                    raise ValueError("table part has an invalid package target")
                if relationship.target in targets:
                    raise ValueError("table package targets must be unique")
                targets.add(relationship.target)
                table = Table.from_tree(fromstring(workbook._archive.read(relationship.target)))
                if not table.name or table.name in names or table.id in identifiers_seen:
                    raise ValueError("table names and identifiers must be nonempty and unique")
                names.add(table.name)
                identifiers_seen.add(table.id)
                tables.add(table)
        except (KeyError, TypeError, ValueError, ET.ParseError) as exc:
            raise ParisonError(f"XLSX workbook has invalid table metadata: {exc}") from exc
        if table := tables.get(name=table_name):
            matches.append((worksheet.title, table))
    if not matches:
        raise ParisonError(f"XLSX table does not exist: {table_name}")
    if len(matches) != 1:
        raise ParisonError(f"XLSX table name is ambiguous: {table_name}")
    worksheet_name, table = matches[0]
    if table.headerRowCount != 1 or table.totalsRowCount or table.totalsRowShown:
        raise ParisonError(f"XLSX table must have one header row and no totals row: {table_name}")
    if any(column.calculatedColumnFormula or column.totalsRowFormula for column in table.tableColumns):
        raise ParisonError(f"XLSX table contains a formula declaration: {table_name}")
    return worksheet_name, table


def _xlsx_defined_name_declarations(path: Path) -> list[tuple[str, int | None]]:
    from openpyxl.xml.functions import iterparse

    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    declarations = []
    try:
        with zipfile.ZipFile(path) as archive:
            with archive.open("xl/workbook.xml") as source:
                overlap = b""
                while chunk := source.read(64 * 1024):
                    if b"\0" in chunk:
                        raise ValueError("workbook metadata must use an ASCII-compatible XML encoding")
                    inspected = (overlap + chunk).upper()
                    if b"<!DOCTYPE" in inspected or b"<!ENTITY" in inspected:
                        raise ValueError("DTD and entity declarations are forbidden")
                    overlap = inspected[-16:]
            with archive.open("xl/workbook.xml") as source:
                for _, element in iterparse(source, events=("end",)):
                    if element.tag == f"{{{namespace}}}definedName":
                        name = element.attrib.get("name")
                        scope = element.attrib.get("localSheetId")
                        if not name:
                            raise ValueError("names must be nonempty")
                        if scope is not None and re.fullmatch(r"[0-9]+", scope) is None:
                            raise ValueError("local sheet scope is invalid")
                        declarations.append((name, int(scope) if scope is not None else None))
                        if len(declarations) > 10_000:
                            raise ValueError("defined-name metadata exceeds 10000 declarations")
                    element.clear()
    except (KeyError, OSError, TypeError, ValueError, zipfile.BadZipFile, ET.ParseError) as exc:
        raise ParisonError(f"XLSX workbook has invalid defined-name metadata: {exc}") from exc
    return declarations


def _xlsx_defined_name(workbook, defined_name: str, declarations: list[tuple[str, int | None]]) -> tuple[str, str]:
    if any(scope is not None and scope >= len(workbook.sheetnames) for _, scope in declarations):
        raise ParisonError("XLSX workbook has invalid defined-name metadata: local sheet scope is invalid")
    identities = [(name, scope) for name, scope in declarations]
    if len(identities) != len(set(identities)):
        raise ParisonError("XLSX workbook has invalid defined-name metadata: names must be unique within scope")
    global_names = {name for name, scope in declarations if scope is None}
    local_names = {name for name, scope in declarations if scope is not None}
    if global_names & local_names:
        raise ParisonError("XLSX workbook has invalid defined-name metadata: global and local names conflict")
    if defined_name not in global_names:
        if defined_name in local_names:
            raise ParisonError(f"XLSX defined name must be workbook-scoped: {defined_name}")
        raise ParisonError(f"XLSX defined name does not exist: {defined_name}")
    definition = workbook.defined_names.get(defined_name)
    if definition is None or definition.localSheetId is not None or definition.type != "RANGE":
        raise ParisonError(f"XLSX defined name must be one static workbook range: {defined_name}")
    try:
        destinations = list(definition.destinations)
    except (AttributeError, TypeError, ValueError) as exc:
        raise ParisonError(f"XLSX defined name has an invalid destination: {defined_name}") from exc
    if len(destinations) != 1:
        raise ParisonError(f"XLSX defined name must resolve to one destination: {defined_name}")
    worksheet_name, cell_range = destinations[0]
    if worksheet_name not in workbook.sheetnames:
        raise ParisonError(f"XLSX defined name must reference one local worksheet: {defined_name}")
    if not re.fullmatch(r"\$[A-Z]+\$[1-9][0-9]*:\$[A-Z]+\$[1-9][0-9]*", cell_range):
        raise ParisonError(f"XLSX defined name must use one absolute bounded rectangle: {defined_name}")
    return worksheet_name, cell_range


def _xlsx_range(cell_range: str, table: bool = False):
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    try:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        if (
            min_col is None
            or min_row is None
            or max_col is None
            or max_row is None
            or min_col > max_col
            or min_row > max_row
            or (table and min_row == max_row)
        ):
            raise ValueError("range is not a bounded header plus data rectangle")
    except ValueError as exc:
        raise ParisonError(f"XLSX range must be one bounded rectangle with a header and data: {cell_range}") from exc
    normalized = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
    return normalized, min_col, min_row, max_col, max_row


def _xlsx_resolved_metadata(source: str | Path) -> dict[str, str]:
    path, kind, worksheet_name, selection = _xlsx_source(source) or (None, None, None, None)
    assert path is not None and selection is not None
    if kind == "range":
        return {"worksheet": worksheet_name, "range": selection}
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ParisonError("XLSX support requires: pip install 'parison[xlsx]'") from exc
    workbook = None
    try:
        declarations = _xlsx_defined_name_declarations(path) if kind == "defined-name" else None
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        if kind == "table":
            worksheet_name, table = _xlsx_table(workbook, selection)
            selected_range = table.ref
            metadata = {"table": selection}
        else:
            assert declarations is not None
            worksheet_name, selected_range = _xlsx_defined_name(workbook, selection, declarations)
            metadata = {"defined_name": selection}
        normalized, *_ = _xlsx_range(selected_range, table=True)
        return {**metadata, "worksheet": worksheet_name, "range": normalized}
    except ParisonError:
        raise
    except Exception as exc:
        raise ParisonError(f"cannot read XLSX selection {selection} from {path}: {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()


def _json_source(source: str | Path) -> tuple[Path, str] | None:
    text = str(source)
    if not text.startswith("json:"):
        return None
    locator = text.removeprefix("json:")
    if "?" in locator or "#" not in locator:
        raise ParisonError("JSON locator must be json:path#pointer without parameters")
    path_text, pointer = locator.split("#", 1)
    if not path_text or (pointer and not pointer.startswith("/")):
        raise ParisonError("JSON locator requires a .json path and an RFC 6901 pointer")
    if Path(path_text).suffix.lower() != ".json":
        raise ParisonError("JSON locator requires a .json document")
    return Path(path_text), pointer


def _xml_source(source: str | Path) -> tuple[Path, str, dict[str, str]] | None:
    text = str(source)
    if not text.startswith("xml:"):
        return None
    locator = text.removeprefix("xml:")
    if "#" not in locator:
        raise ParisonError("XML locator must be xml:path#record-path with optional namespace bindings")
    path_text, selection = locator.rsplit("#", 1)
    record_path, separator, query = selection.partition("?")
    if not path_text or not record_path or Path(path_text).suffix.lower() != ".xml":
        raise ParisonError("XML locator requires a .xml document and bounded record path")
    segments = record_path.split("/")
    if not 1 <= len(segments) <= 32 or any(
        not re.fullmatch(r"(?:[A-Za-z_][\w.-]*:)?[A-Za-z_][\w.-]*", segment, re.ASCII) for segment in segments
    ):
        raise ParisonError("XML record path must contain 1 to 32 bounded QName segments")
    namespaces = {}
    if separator:
        try:
            bindings = parse_qsl(query, keep_blank_values=True, strict_parsing=True)
        except ValueError as exc:
            raise ParisonError("XML namespace bindings are invalid") from exc
        for name, binding in bindings:
            if name != "ns" or "," not in binding:
                raise ParisonError("XML locator parameters must be ns=prefix,URI bindings")
            prefix, uri = binding.split(",", 1)
            if (
                not re.fullmatch(r"[A-Za-z_][\w.-]*", prefix, re.ASCII)
                or not uri
                or prefix in namespaces
                or uri in namespaces.values()
            ):
                raise ParisonError("XML namespace prefixes and URIs must be nonempty and unique")
            namespaces[prefix] = uri
    used_prefixes = {segment.split(":", 1)[0] for segment in segments if ":" in segment}
    if not used_prefixes <= set(namespaces):
        raise ParisonError("XML record path prefixes must be declared")
    return Path(path_text), record_path, namespaces


def _source_path(source: str | Path) -> Path:
    sqlite_source = _sqlite_source(source)
    xlsx_source = _xlsx_source(source)
    json_source = _json_source(source)
    xml_source = _xml_source(source)
    return (
        sqlite_source[0]
        if sqlite_source
        else xlsx_source[0]
        if xlsx_source
        else json_source[0]
        if json_source
        else xml_source[0]
        if xml_source
        else Path(source)
    )


def _source_paths(source: str | Path) -> tuple[Path, ...]:
    path = _source_path(source)
    if (
        _sqlite_source(source)
        or _xlsx_source(source)
        or _json_source(source)
        or _xml_source(source)
        or not path.is_dir()
    ):
        return (path,)
    if path.is_symlink():
        raise ParisonError(f"partitioned input must not be a symlink: {path}")
    try:
        paths = tuple(
            sorted((item for item in path.iterdir() if not item.name.startswith(".")), key=lambda item: item.name)
        )
    except OSError as exc:
        raise ParisonError(f"cannot inspect partitioned input {path}: {exc}") from exc
    if not paths:
        raise ParisonError(f"partitioned input has no files: {path}")
    if any(item.is_symlink() or not item.is_file() for item in paths):
        raise ParisonError(f"partitioned input must contain only regular non-symlink files: {path}")
    formats = {_file_format(item) for item in paths}
    if None in formats or len(formats) != 1:
        raise ParisonError(f"partitioned input must contain one supported file format: {path}")
    if formats == {"xlsx"}:
        raise ParisonError(f"XLSX partitioned inputs are not supported: {path}")
    return paths


def _source_digest(source: str | Path) -> str:
    paths = _source_paths(source)
    if len(paths) == 1 and not _source_path(source).is_dir():
        return _digest(paths[0])
    digest = hashlib.sha256()
    for path in paths:
        digest.update(path.name.encode("utf-8"))
        digest.update(b"\0")
        digest.update(bytes.fromhex(_digest(path)))
    return digest.hexdigest()


def _source_bytes(source: str | Path) -> int:
    try:
        return sum(path.stat().st_size for path in _source_paths(source))
    except OSError as exc:
        raise ParisonError(f"cannot inspect inputs: {exc}") from exc


def _decoded_sizes(sources: tuple[str | Path, ...], max_decoded_bytes: int) -> dict[Path, int]:
    if max_decoded_bytes <= 0:
        raise ParisonError("max_decoded_bytes must be positive")
    total = 0
    sizes = {}
    for source in sources:
        for path in _source_paths(source):
            if path.suffix.lower() == ".xlsx":
                if path.is_symlink():
                    raise ParisonError(f"XLSX workbook must be a regular non-symlink file: {path}")
                try:
                    with zipfile.ZipFile(path) as archive:
                        members = archive.infolist()
                        if len(members) > 10_000:
                            raise ParisonError(f"XLSX archive has too many members: {path}")
                        if len({member.filename for member in members}) != len(members):
                            raise ParisonError(f"XLSX archive contains duplicate member names: {path}")
                        if any(member.flag_bits & 1 for member in members):
                            raise ParisonError(f"XLSX archive contains encrypted members: {path}")
                        size = sum(member.file_size for member in members)
                except ParisonError:
                    raise
                except (OSError, zipfile.BadZipFile) as exc:
                    raise ParisonError(f"cannot inspect XLSX archive {path}: {exc}") from exc
                total += size
                if total > max_decoded_bytes:
                    raise ParisonError(f"combined decoded input size exceeds limit {max_decoded_bytes} bytes")
                sizes[path] = size
                continue
            if path.suffix.lower() != ".gz":
                continue
            size = 0
            try:
                with gzip.open(path, "rb") as handle:
                    while chunk := handle.read(min(1024 * 1024, max_decoded_bytes - total + 1)):
                        size += len(chunk)
                        total += len(chunk)
                        if total > max_decoded_bytes:
                            raise ParisonError(f"combined decoded input size exceeds limit {max_decoded_bytes} bytes")
            except ParisonError:
                raise
            except (OSError, EOFError, gzip.BadGzipFile) as exc:
                raise ParisonError(f"cannot decompress {path}: {exc}") from exc
            sizes[path] = size
    return sizes


def _source_metadata(
    source: str | Path,
    digest: str,
    decoded_sizes: dict[Path, int] | None = None,
    delimiter: str = ",",
    null_tokens: list[str] | None = None,
    include_xlsx_table_selection: bool = True,
) -> dict[str, Any]:
    path = _source_path(source)
    paths = _source_paths(source)
    format_name = (
        "sqlite"
        if _sqlite_source(source)
        else "xlsx"
        if _xlsx_source(source)
        else "json"
        if _json_source(source)
        else "xml"
        if _xml_source(source)
        else (_file_format(paths[0]) or path.suffix.lower().lstrip("."))
    )
    metadata = {"sha256": digest, "bytes": _source_bytes(source), "format": format_name}
    if format_name in {"csv", "tsv"}:
        metadata["delimiter"] = delimiter
        metadata["null_tokens"] = null_tokens or []
    sqlite_source = _sqlite_source(source)
    if sqlite_source:
        metadata.update(format="sqlite", table=sqlite_source[1])
    elif _xlsx_source(source):
        selection = _xlsx_resolved_metadata(source)
        if include_xlsx_table_selection or not ({"table", "defined_name"} & set(selection)):
            metadata.update(selection)
        metadata["format"] = "xlsx"
    elif json_source := _json_source(source):
        metadata.update(format="json", pointer=json_source[1])
    elif xml_source := _xml_source(source):
        metadata.update(format="xml", record_path=xml_source[1], namespaces=dict(sorted(xml_source[2].items())))
    elif path.is_dir():
        metadata["partitions"] = len(paths)
    if decoded_sizes and any(item in decoded_sizes for item in paths):
        metadata.update(
            compression="zip" if format_name == "xlsx" else "gzip",
            decoded_bytes=sum(decoded_sizes.get(item, 0) for item in paths),
        )
    return metadata


def _xlsx_rows(source: str | Path, max_rows: int | None):
    path, kind, worksheet_name, selection = _xlsx_source(source) or (None, None, None, None)
    assert path is not None and selection is not None
    if path.is_symlink() or not path.is_file():
        raise ParisonError(f"XLSX workbook must be a regular non-symlink file: {path}")
    try:
        from openpyxl import load_workbook
        from openpyxl.styles.numbers import is_datetime
    except ImportError as exc:
        raise ParisonError("XLSX support requires: pip install 'parison[xlsx]'") from exc
    workbook = None
    try:
        declarations = _xlsx_defined_name_declarations(path) if kind == "defined-name" else None
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        table = None
        if kind == "table":
            worksheet_name, table = _xlsx_table(workbook, selection)
            cell_range, min_col, min_row, max_col, max_row = _xlsx_range(table.ref, table=True)
        elif kind == "defined-name":
            assert declarations is not None
            worksheet_name, defined_range = _xlsx_defined_name(workbook, selection, declarations)
            cell_range, min_col, min_row, max_col, max_row = _xlsx_range(defined_range, table=True)
        else:
            cell_range = selection
            _, min_col, min_row, max_col, max_row = _xlsx_range(cell_range)
        if worksheet_name not in workbook.sheetnames:
            raise ParisonError(f"XLSX worksheet does not exist: {worksheet_name}")
        if max_rows is not None and max_row - min_row > max_rows:
            raise ParisonError(f"row count in {source} exceeds limit {max_rows}")
        worksheet = workbook[worksheet_name]
        rows = worksheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col)
        header_cells = next(rows)
        if any(cell.data_type == "f" for cell in header_cells):
            raise ParisonError(f"XLSX range contains a formula cell: {cell_range}")
        if any(cell.data_type == "e" for cell in header_cells):
            raise ParisonError(f"XLSX range contains an error cell: {cell_range}")
        headers = [cell.value for cell in header_cells]
        if any(not isinstance(name, str) or not name for name in headers) or len(headers) != len(set(headers)):
            raise ParisonError(f"XLSX range has empty, non-string, or duplicate column names: {cell_range}")
        if table is not None and headers != table.column_names:
            raise ParisonError(f"XLSX table columns do not match stored header cells: {selection}")
        yield headers, None
        for cells in rows:
            if any(cell.data_type == "f" for cell in cells):
                raise ParisonError(f"XLSX range contains a formula cell: {cell_range}")
            if any(cell.data_type == "e" for cell in cells):
                raise ParisonError(f"XLSX range contains an error cell: {cell_range}")
            values = [
                cell.value.date()
                if type(cell.value) is datetime and is_datetime(cell.number_format) == "date"
                else cell.value
                for cell in cells
            ]
            yield headers, dict(zip(headers, values, strict=True))
    except ParisonError:
        raise
    except Exception as exc:
        raise ParisonError(f"cannot read XLSX selection from {path}: {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()


def _sqlite_rows(source: str | Path, max_rows: int):
    path, table = _sqlite_source(source) or (None, None)
    assert path is not None and table is not None
    if path.is_symlink() or not path.is_file():
        raise ParisonError(f"SQLite database must be a regular non-symlink file: {path}")
    if any(Path(str(path) + suffix).exists() for suffix in ("-wal", "-journal")):
        raise ParisonError(f"SQLite database has an active journal sidecar: {path}")
    cursors = []
    try:
        connection = sqlite3.connect(path.resolve().as_uri() + "?mode=ro", uri=True)
        cursors.append(connection.execute("PRAGMA query_only=ON"))
        lookup = connection.execute("SELECT type FROM sqlite_schema WHERE name=?", (table,))
        cursors.append(lookup)
        found = lookup.fetchone()
        if found != ("table",):
            raise ParisonError(f"SQLite object is not an ordinary table: {table}")
        cursor = connection.execute(f'SELECT * FROM "{table.replace(chr(34), chr(34) * 2)}"')
        cursors.append(cursor)
        headers = [item[0] for item in cursor.description or []]
        if not headers or any(not name for name in headers) or len(headers) != len(set(headers)):
            raise ParisonError(f"SQLite table has empty or duplicate columns: {table}")
        yield headers, None
        for row_count, values in enumerate(cursor):
            if row_count >= max_rows:
                raise ParisonError(f"row count in {source} exceeds limit {max_rows}")
            if any(isinstance(value, bytes) for value in values):
                raise ParisonError(f"SQLite table contains an unsupported BLOB value: {table}")
            yield headers, dict(zip(headers, values, strict=False))
    except ParisonError:
        raise
    except sqlite3.Error as exc:
        raise ParisonError(f"cannot read SQLite table {table} from {path}: {exc}") from exc
    finally:
        for opened_cursor in reversed(cursors):
            opened_cursor.close()
        if "connection" in locals():
            connection.close()


def _jsonl_record(line: str, path: Path, line_number: int) -> dict[str, Any]:
    if not line.strip():
        raise ParisonError(f"blank JSON Lines record at line {line_number} in {path}")

    def object_pairs(pairs):
        value = dict(pairs)
        if len(value) != len(pairs):
            raise ParisonError(f"duplicate JSON key at line {line_number} in {path}")
        return value

    try:
        value = json.loads(
            line,
            object_pairs_hook=object_pairs,
            parse_int=str,
            parse_float=str,
            parse_constant=lambda token: (_ for _ in ()).throw(ValueError(token)),
        )
    except (json.JSONDecodeError, ValueError) as exc:
        raise ParisonError(f"cannot parse JSON at line {line_number} in {path}: {exc}") from exc
    if not isinstance(value, dict):
        raise ParisonError(f"JSON Lines record {line_number} in {path} must be an object")
    if any(not isinstance(name, str) or not name for name in value):
        raise ParisonError(f"JSON Lines record {line_number} in {path} has an empty or invalid field name")
    if any(isinstance(item, (dict, list)) for item in value.values()):
        raise ParisonError(f"JSON Lines record {line_number} in {path} contains a nested value")
    return {name: ("true" if item is True else "false" if item is False else item) for name, item in value.items()}


def _resolve_json_pointer(value: Any, pointer: str, where: str) -> Any:
    failure = f"JSON pointer does not resolve{f' {where}' if where else ''}: {pointer}"
    for raw_token in pointer.split("/")[1:] if pointer else ():
        if "~" in raw_token.replace("~0", "").replace("~1", ""):
            raise ParisonError(f"JSON pointer has an invalid escape: {pointer}")
        token = raw_token.replace("~1", "/").replace("~0", "~")
        if isinstance(value, dict) and token in value:
            value = value[token]
        elif isinstance(value, list) and (
            token == "0"
            or (token and token[0] in "123456789" and all(character in "0123456789" for character in token))
        ):
            try:
                index = int(token)
            except ValueError as exc:
                raise ParisonError(failure) from exc
            if index >= len(value):
                raise ParisonError(failure)
            value = value[index]
        else:
            raise ParisonError(failure)
    return value


def _json_rows(source: str | Path, max_rows: int | None, selectors: dict[str, str] | None = None):
    path, pointer = _json_source(source) or (None, None)
    assert path is not None and pointer is not None
    if path.is_symlink() or not path.is_file():
        raise ParisonError(f"JSON document must be a regular non-symlink file: {path}")

    def object_pairs(pairs):
        value = dict(pairs)
        if len(value) != len(pairs):
            raise ParisonError(f"duplicate JSON key in {path}")
        return value

    try:
        value = json.loads(
            path.read_text(encoding="utf-8"),
            object_pairs_hook=object_pairs,
            parse_int=str,
            parse_float=str,
            parse_constant=lambda token: (_ for _ in ()).throw(ValueError(token)),
        )
    except ParisonError:
        raise
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError, RecursionError) as exc:
        raise ParisonError(f"cannot parse JSON document {path}: {exc}") from exc
    value = _resolve_json_pointer(value, pointer, "")
    if not isinstance(value, list):
        raise ParisonError(f"JSON pointer must select an array: {pointer}")
    if max_rows is not None and len(value) > max_rows:
        raise ParisonError(f"row count in {source} exceeds limit {max_rows}")
    headers = None
    for index, row in enumerate(value):
        if not isinstance(row, dict):
            raise ParisonError(f"JSON record {index} in {path} must be an object")
        if any(not isinstance(name, str) or not name for name in row):
            raise ParisonError(f"JSON record {index} in {path} has an empty or invalid field name")
        if selectors is None:
            if any(isinstance(item, (dict, list)) for item in row.values()):
                raise ParisonError(f"JSON record {index} in {path} contains a nested value")
            projected = row
        else:
            projected = {}
            for name, field_pointer in selectors.items():
                item = _resolve_json_pointer(row, field_pointer, f"at record {index}")
                if isinstance(item, (dict, list)):
                    raise ParisonError(f"JSON field pointer for {name} selects a container at record {index} in {path}")
                projected[name] = item
        if headers is None:
            headers = list(projected)
            if not headers:
                raise ParisonError(f"JSON record {index} in {path} has no fields")
            yield headers, None
        elif selectors is None and set(projected) != set(headers):
            raise ParisonError(f"inconsistent JSON fields at record {index} in {path}")
        yield (
            headers,
            {name: "true" if item is True else "false" if item is False else item for name, item in projected.items()},
        )
    if headers is None:
        raise ParisonError(f"input has no schema: {source}")


def _xml_rows(source: str | Path, max_rows: int | None, field_projection: dict[str, Any] | None = None):
    path, record_path, namespaces = _xml_source(source) or (None, None, None)
    assert path is not None and record_path is not None and namespaces is not None
    if path.is_symlink() or not path.is_file():
        raise ParisonError(f"XML document must be a regular non-symlink file: {path}")
    try:
        content = path.read_bytes()
    except OSError as exc:
        raise ParisonError(f"cannot read XML document {path}: {exc}") from exc
    if len(content) > 64 * 1024 * 1024:
        raise ParisonError(f"XML document exceeds 67108864 byte parser limit: {path}")
    try:
        xml_text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ParisonError(f"XML document must use UTF-8 encoding: {path}") from exc
    declaration = re.match(r"""<\?xml\s+[^?]*\bencoding\s*=\s*(["'])([^"']+)\1""", xml_text, re.IGNORECASE)
    if declaration and declaration.group(2).lower().replace("_", "-") not in {"utf-8", "utf8"}:
        raise ParisonError(f"XML document must use UTF-8 encoding: {path}")
    offset = 0
    while (declaration := xml_text.find("<!", offset)) >= 0:
        if xml_text.startswith("<!--", declaration):
            end = xml_text.find("-->", declaration + 4)
            offset = len(xml_text) if end < 0 else end + 3
        elif xml_text.startswith("<![CDATA[", declaration):
            end = xml_text.find("]]>", declaration + 9)
            offset = len(xml_text) if end < 0 else end + 3
        elif xml_text.startswith(("<!DOCTYPE", "<!ENTITY"), declaration):
            break
        else:
            offset = declaration + 2
    else:
        declaration = -1
    if declaration >= 0 and xml_text.startswith(("<!DOCTYPE", "<!ENTITY"), declaration):
        raise ParisonError(f"XML document contains a forbidden DTD or entity declaration: {path}")
    reverse_namespaces = {uri: prefix for prefix, uri in namespaces.items()}

    def expanded(name: str) -> str:
        if ":" not in name:
            return name
        prefix, local = name.split(":", 1)
        return f"{{{namespaces[prefix]}}}{local}"

    def projected(name: str) -> str:
        if not name.startswith("{"):
            return name
        uri, local = name[1:].split("}", 1)
        if uri not in reverse_namespaces:
            raise ParisonError(f"XML field uses an undeclared namespace: {uri}")
        return f"{reverse_namespaces[uri]}:{local}"

    expected_path = tuple(expanded(segment) for segment in record_path.split("/"))
    selector_namespaces = field_projection.get("namespaces", {}) if field_projection else {}

    def expanded_selector(name: str) -> str:
        if ":" not in name:
            return name
        prefix, local = name.split(":", 1)
        return f"{{{selector_namespaces[prefix]}}}{local}"

    attribute_fields = {}
    element_trie: dict[str, Any] = {}
    if field_projection:
        for name, selector in field_projection["fields"].items():
            if selector.startswith("@"):
                attribute_fields[expanded_selector(selector[1:])] = name
                continue
            node = element_trie
            for segment in selector.split("/"):
                node = node.setdefault(expanded_selector(segment), {})
            node[None] = name

    def projected_record(element: ET.Element) -> dict[str, Any]:
        if set(element.attrib) != set(attribute_fields):
            raise ParisonError(f"XML record attributes do not match field projection at {record_path}")
        if (element.text or "").strip():
            raise ParisonError(f"XML record has mixed text at {record_path}")
        if element.tail and element.tail.strip():
            raise ParisonError(f"XML record has mixed tail text at {record_path}")
        row = {name: element.attrib[tag] for tag, name in attribute_fields.items()}

        def walk(parent: ET.Element, trie: dict[str | None, Any]) -> None:
            expected_children = {tag for tag in trie if tag is not None}
            children: dict[str, ET.Element] = {}
            for child in parent:
                if child.tag not in expected_children:
                    raise ParisonError(f"XML record contains an unselected element at {record_path}")
                if child.tag in children:
                    raise ParisonError(f"XML field path is repeated at {record_path}")
                if child.tail and child.tail.strip():
                    raise ParisonError(f"XML field contains mixed tail text at {record_path}")
                children[child.tag] = child
            if set(children) != expected_children:
                raise ParisonError(f"XML field path does not resolve at {record_path}")
            for tag, child in children.items():
                branch = trie[tag]
                if None in branch:
                    if child.attrib or len(child):
                        raise ParisonError(f"XML field must be an attribute-free text element at {record_path}")
                    row[branch[None]] = child.text or ""
                else:
                    if child.attrib or (child.text or "").strip():
                        raise ParisonError("XML intermediate field element must contain only selected children")
                    walk(child, branch)

        walk(element, element_trie)
        return row

    headers = None
    rows = 0
    elements = 0
    path_stack = []
    element_stack = []
    try:
        for event, element in ET.iterparse(io.BytesIO(content), events=("start", "end")):
            if event == "start":
                elements += 1
                path_stack.append(element.tag)
                element_stack.append(element)
                if elements > 1_000_000:
                    raise ParisonError(f"XML document exceeds 1000000 element limit: {path}")
                if len(path_stack) > 100:
                    raise ParisonError(f"XML document exceeds depth limit 100: {path}")
                continue
            if element.text is not None and len(element.text.encode("utf-8")) > 1024 * 1024:
                raise ParisonError(f"XML element text exceeds 1048576 byte limit: {path}")
            if element.tail is not None and len(element.tail.encode("utf-8")) > 1024 * 1024:
                raise ParisonError(f"XML element text exceeds 1048576 byte limit: {path}")
            if tuple(path_stack) == expected_path:
                if max_rows is not None and rows >= max_rows:
                    raise ParisonError(f"row count in {source} exceeds limit {max_rows}")
                if field_projection:
                    row = projected_record(element)
                else:
                    if element.attrib or (element.text or "").strip():
                        raise ParisonError(f"XML record has attributes or mixed text at {record_path}")
                    row = {}
                    for child in element:
                        if child.attrib or len(child) or (child.tail and child.tail.strip()):
                            raise ParisonError(f"XML field must be an attribute-free text element at {record_path}")
                        name = projected(child.tag)
                        if name in row:
                            raise ParisonError(f"XML record has a duplicate field {name} at {record_path}")
                        row[name] = child.text or ""
                if not row:
                    raise ParisonError(f"XML record at {record_path} has no fields")
                if headers is None:
                    headers = list(row)
                    yield headers, None
                elif set(row) != set(headers):
                    raise ParisonError(f"inconsistent XML fields at record {rows} in {path}")
                rows += 1
                yield headers, row
            preserve_for_record = (
                len(path_stack) > len(expected_path) and tuple(path_stack[: len(expected_path)]) == expected_path
            )
            parent = element_stack[-2] if len(element_stack) > 1 else None
            if parent is not None and not preserve_for_record:
                parent.remove(element)
            path_stack.pop()
            element_stack.pop()
    except ParisonError:
        raise
    except (ET.ParseError, RecursionError) as exc:
        raise ParisonError(f"cannot parse XML document {path}: {exc}") from exc
    if headers is None:
        raise ParisonError(f"XML record path selects no elements: {record_path}")


def _iter_jsonl(path: Path, max_rows: int, decoded_sizes: dict[Path, int]):
    try:
        handle = _open_text(path, decoded_sizes)
    except (OSError, EOFError, gzip.BadGzipFile) as exc:
        raise ParisonError(f"cannot read {path}: {exc}") from exc
    with handle:
        headers = None
        try:
            for line_number, line in enumerate(handle, 1):
                if line_number > max_rows:
                    raise ParisonError(f"row count in {path} exceeds limit {max_rows}")
                row = _jsonl_record(line, path, line_number)
                if headers is None:
                    headers = list(row)
                elif set(row) != set(headers):
                    raise ParisonError(f"inconsistent JSON Lines fields at line {line_number} in {path}")
                yield row
        except UnicodeDecodeError as exc:
            raise ParisonError(f"cannot parse {path}: {exc}") from exc
        if headers is None:
            raise ParisonError(f"input has no schema: {path}")


def _read(
    path: str | Path, recipe: dict[str, Any], max_rows: int, side: str, decoded_sizes: dict[Path, int] | None = None
) -> tuple[list[dict[str, Any]], list[str]]:
    raw_rows = list(_iter_input_rows(path, recipe, max_rows, side, decoded_sizes or {}))
    rows = [_parse_row(row, recipe, side, recipe.get("output", {}).get("sensitivity") == "raw") for row in raw_rows]
    return rows, []


def _digest(path: Path) -> str:
    digest = hashlib.sha256()
    try:
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
    except OSError as exc:
        raise ParisonError(f"cannot read {path}: {exc}") from exc
    return digest.hexdigest()


def _validate_parquet_projection(schema: Any, recipe: dict[str, Any], side: str) -> None:
    if recipe.get("input_projection") != "semantic-scalar-v1":
        return
    scalar = {
        "Null",
        "String",
        "Boolean",
        "Int8",
        "Int16",
        "Int32",
        "Int64",
        "UInt8",
        "UInt16",
        "UInt32",
        "UInt64",
        "Float32",
        "Float64",
        "Decimal",
        "Date",
        "Datetime",
    }
    admitted = {
        "string": {"String"},
        "integer": {"Int8", "Int16", "Int32", "Int64", "UInt8", "UInt16", "UInt32", "UInt64", "String"},
        "decimal": {"Int8", "Int16", "Int32", "Int64", "UInt8", "UInt16", "UInt32", "UInt64", "Decimal", "String"},
        "float": {
            "Int8",
            "Int16",
            "Int32",
            "Int64",
            "UInt8",
            "UInt16",
            "UInt32",
            "UInt64",
            "Float32",
            "Float64",
            "String",
        },
        "boolean": {"Boolean", "String"},
        "date": {"Date", "String"},
        "timestamp": {"Datetime", "String"},
    }
    for dtype in schema.values():
        if str(dtype.base_type()) not in scalar:
            raise ParisonError(f"Parquet input has unsupported semantic-scalar-v1 type {dtype}")
    for name, policy in recipe["columns"].items():
        source = _source_name(recipe, name, side)
        if source not in schema:
            continue
        dtype = schema[source]
        base = str(dtype.base_type())
        if base != "Null" and base not in admitted[policy["type"]]:
            raise ParisonError(f"Parquet column {source} cannot project to {policy['type']} under semantic-scalar-v1")
        if base == "Datetime" and (
            getattr(dtype, "time_unit", None) == "ns" or getattr(dtype, "time_zone", None) is None
        ):
            raise ParisonError(
                f"Parquet column {source} has lossy or timezone-naive timestamp type under semantic-scalar-v1"
            )


def _input_columns(
    path: str | Path,
    decoded_sizes: dict[Path, int] | None = None,
    delimiter: str = ",",
    recipe: dict[str, Any] | None = None,
    side: str | None = None,
    max_rows: int | None = None,
) -> list[str]:
    decoded_sizes = decoded_sizes or {}
    source_path = _source_path(path)
    if source_path.is_dir():
        schemas = [
            _input_columns(partition, decoded_sizes, delimiter, recipe, side, max_rows)
            for partition in _source_paths(path)
        ]
        if any(set(schema) != set(schemas[0]) for schema in schemas[1:]):
            raise ParisonError(f"partitioned input has inconsistent schemas: {source_path}")
        return schemas[0]
    if _sqlite_source(path):
        rows = _sqlite_rows(path, 1)
        try:
            columns, _ = next(rows)
            return columns
        finally:
            rows.close()
    if _xlsx_source(path):
        rows = _xlsx_rows(path, None)
        try:
            columns, _ = next(rows)
            return columns
        finally:
            rows.close()
    if _json_source(path):
        projection = recipe.get("field_projection", {}).get(side) if recipe is not None and side is not None else None
        selectors = projection["fields"] if projection and projection["kind"] == "json-pointer-v1" else None
        rows = _json_rows(path, max_rows, selectors)
        try:
            columns, _ = next(rows)
            for _ in rows:
                pass
            return columns
        finally:
            rows.close()
    if _xml_source(path):
        projection = recipe.get("field_projection", {}).get(side) if recipe is not None and side is not None else None
        rows = _xml_rows(path, max_rows, projection if projection and projection["kind"] == "xml-field-v1" else None)
        try:
            columns, _ = next(rows)
            for _ in rows:
                pass
            return columns
        finally:
            rows.close()
    path = Path(path)
    if path.is_symlink():
        raise ParisonError(f"input must not be a symlink: {path}")
    if not path.is_file():
        raise ParisonError(f"input is not a regular file: {path}")
    format_name = _file_format(path)
    if format_name in {"csv", "tsv"}:
        try:
            with _open_text(path, decoded_sizes, newline="") as handle:
                columns = next(csv.reader(handle, delimiter=delimiter, strict=True), [])
        except (OSError, EOFError, gzip.BadGzipFile, csv.Error, UnicodeDecodeError) as exc:
            raise ParisonError(f"cannot read schema from {path}: {exc}") from exc
    elif format_name == "jsonl":
        try:
            with _open_text(path, decoded_sizes) as handle:
                line = next(handle, "")
            columns = list(_jsonl_record(line, path, 1)) if line else []
        except (OSError, EOFError, gzip.BadGzipFile, UnicodeDecodeError) as exc:
            raise ParisonError(f"cannot read schema from {path}: {exc}") from exc
    elif format_name == "parquet":
        try:
            import polars as pl
        except ImportError as exc:
            raise ParisonError("Parquet support requires: pip install 'parison[parquet]'") from exc
        try:
            schema = pl.scan_parquet(path).collect_schema()
            if recipe is not None and side is not None:
                _validate_parquet_projection(schema, recipe, side)
            columns = schema.names()
        except Exception as exc:
            raise ParisonError(f"cannot read schema from {path}: {exc}") from exc
    else:
        raise ParisonError(
            f"unsupported input format for {path}; use .csv, .tsv, .jsonl, .parquet, an XLSX locator, "
            "or gzip-compressed text"
        )
    if not columns:
        raise ParisonError(f"input has no schema: {path}")
    if any(not isinstance(name, str) or not name for name in columns) or len(columns) != len(set(columns)):
        raise ParisonError(f"input has empty or duplicate column names: {path}")
    return columns


def _validate_field_projection_source(path: str | Path, recipe: dict[str, Any], side: str) -> None:
    projection = recipe.get("field_projection", {}).get(side)
    if projection is None:
        return
    json_source = _json_source(path)
    xml_source = _xml_source(path)
    if json_source and projection["kind"] != "json-pointer-v1":
        raise ParisonError(f"field_projection.{side} for a JSON document must use json-pointer-v1")
    if not json_source and projection["kind"] == "json-pointer-v1":
        raise ParisonError(f"field_projection.{side} kind json-pointer-v1 requires a JSON document locator")
    if xml_source and projection["kind"] != "xml-field-v1":
        raise ParisonError(f"field_projection.{side} for an XML document must use xml-field-v1")
    if not xml_source and projection["kind"] == "xml-field-v1":
        raise ParisonError(f"field_projection.{side} kind xml-field-v1 requires an XML document locator")


def _record_diagnostics(
    source: str | Path,
    recipe: dict[str, Any],
    max_rows: int,
    side: str,
    decoded_sizes: dict[Path, int],
) -> dict[str, Any]:
    seen: set[tuple[Any, ...]] = set()
    duplicates: set[tuple[Any, ...]] = set()
    invalid_fields: Counter[str] = Counter()
    invalid_rows = null_key_rows = rows = 0
    for raw in _iter_input_rows(source, recipe, max_rows, side, decoded_sizes):
        rows += 1
        row = {}
        row_is_invalid = False
        for name, policy in recipe["columns"].items():
            try:
                parsed = _parse_input(raw.get(_source_name(recipe, name, side)), policy, name, recipe)
                row[name] = _normalize(parsed, policy)
            except ParisonError:
                invalid_fields[name] += 1
                row_is_invalid = True
        invalid_rows += row_is_invalid
        if not set(recipe["keys"]) <= set(row):
            continue
        key = tuple(row[name] for name in recipe["keys"])
        null_key_rows += any(value is None for value in key)
        if key in seen:
            duplicates.add(key)
        seen.add(key)
    return {
        "status": "invalid" if invalid_rows or null_key_rows or duplicates else "valid",
        "rows": rows,
        "invalid_rows": invalid_rows,
        "invalid_fields": dict(sorted(invalid_fields.items())),
        "null_key_rows": null_key_rows,
        "duplicate_keys": len(duplicates),
    }


def _aggregate_record_diagnostics(
    source: str | Path,
    recipe: dict[str, Any],
    max_rows: int,
    max_groups: int,
    side: str,
    decoded_sizes: dict[Path, int],
) -> dict[str, Any]:
    groups: set[tuple[Any, ...]] = set()
    invalid_fields: Counter[str] = Counter()
    invalid_rows = null_group_rows = rejected_null_values = rows = 0
    group_limit_exceeded = False
    rejected_columns = {
        policy["column"]
        for policy in recipe["measures"].values()
        if policy["operator"] != "count" and policy["nulls"] == "reject"
    }
    for raw in _iter_input_rows(source, recipe, max_rows, side, decoded_sizes):
        rows += 1
        parsed = {}
        row_is_invalid = False
        for name, policy in recipe["columns"].items():
            try:
                parsed[name] = _normalize(
                    _parse_input(raw.get(_source_name(recipe, name, side)), policy, name, recipe), policy
                )
            except ParisonError:
                invalid_fields[name] += 1
                row_is_invalid = True
        invalid_rows += row_is_invalid
        if set(recipe["group_by"]) <= set(parsed):
            key = tuple(parsed[name] for name in recipe["group_by"])
            if any(value is None for value in key):
                null_group_rows += 1
            elif len(groups) < max_groups or key in groups:
                groups.add(key)
            else:
                group_limit_exceeded = True
        rejected_null_values += sum(parsed.get(name) is None for name in rejected_columns if name in parsed)
    invalid = invalid_rows or null_group_rows or rejected_null_values or group_limit_exceeded
    return {
        "status": "invalid" if invalid else "valid",
        "rows": rows,
        "invalid_rows": invalid_rows,
        "invalid_fields": dict(sorted(invalid_fields.items())),
        "null_group_rows": null_group_rows,
        "rejected_null_measure_values": rejected_null_values,
        "groups": len(groups),
        "group_limit_exceeded": group_limit_exceeded,
    }


def _multiset_record_diagnostics(source, recipe, max_rows, max_distinct_rows, side, decoded_sizes):
    seen: set[tuple[Any, ...]] = set()
    invalid_fields: Counter[str] = Counter()
    invalid_rows = rows = 0
    for raw in _iter_input_rows(source, recipe, max_rows, side, decoded_sizes):
        rows += 1
        parsed = {}
        row_invalid = False
        for name, policy in recipe["columns"].items():
            try:
                parsed[name] = _normalize(
                    _parse_input(raw.get(_source_name(recipe, name, side)), policy, name, recipe), policy
                )
            except ParisonError:
                invalid_fields[name] += 1
                row_invalid = True
        invalid_rows += row_invalid
        if not row_invalid:
            seen.add(tuple(parsed[name] for name in sorted(recipe["columns"])))
    limit_exceeded = len(seen) > max_distinct_rows
    return {
        "status": "invalid" if invalid_rows or limit_exceeded else "valid",
        "rows": rows,
        "invalid_rows": invalid_rows,
        "invalid_fields": dict(sorted(invalid_fields.items())),
        "distinct_rows": len(seen),
        "distinct_row_limit_exceeded": limit_exceeded,
    }


def validate_inputs(
    recipe_path: str | Path,
    baseline: str | Path,
    candidate: str | Path,
    max_input_bytes: int = 1_000_000_000,
    expected_policy_sha256: str | None = None,
    max_decoded_bytes: int = 1_000_000_000,
    validate_records: bool = False,
    max_rows: int = 5_000_000,
    max_groups: int = 100_000,
    max_distinct_rows: int = 100_000,
) -> dict[str, Any]:
    """Validate input schemas against a recipe without comparing records."""
    if max_input_bytes <= 0:
        raise ParisonError("max_input_bytes must be positive")
    if max_rows <= 0:
        raise ParisonError("max_rows must be positive")
    if max_groups <= 0:
        raise ParisonError("max_groups must be positive")
    if max_distinct_rows <= 0:
        raise ParisonError("max_distinct_rows must be positive")
    recipe = load_recipe(recipe_path)
    policy_sha256 = _checked_policy_sha256(recipe, expected_policy_sha256)
    sizes = {side: _source_bytes(source) for side, source in (("baseline", baseline), ("candidate", candidate))}
    if sum(sizes.values()) > max_input_bytes:
        raise ParisonError(f"combined input size {sum(sizes.values())} exceeds limit {max_input_bytes} bytes")
    decoded_sizes = _decoded_sizes((baseline, candidate), max_decoded_bytes)
    before = {str(source): _source_digest(source) for source in (baseline, candidate)} if validate_records else {}
    inputs = {}
    status = "valid"
    for side, source in (("baseline", baseline), ("candidate", candidate)):
        delimiter = recipe["delimiters"][side]
        _validate_field_projection_source(source, recipe, side)
        if _xlsx_source(source) and recipe.get("input_projection") != "semantic-scalar-v1":
            raise ParisonError("XLSX inputs require input_projection 'semantic-scalar-v1'")
        if _json_source(source) and recipe.get("input_projection") != "semantic-scalar-v1":
            raise ParisonError("JSON document inputs require input_projection 'semantic-scalar-v1'")
        if _xml_source(source) and recipe.get("input_projection") != "semantic-scalar-v1":
            raise ParisonError("XML inputs require input_projection 'semantic-scalar-v1'")
        columns = _input_columns(source, decoded_sizes, delimiter, recipe, side, max_rows)
        schema = _schema_details(columns, recipe, side)
        if schema["missing"] or schema["unexpected"]:
            status = "invalid"
        paths = _source_paths(source)
        sqlite_source = _sqlite_source(source)
        xlsx_source = _xlsx_source(source)
        json_source = _json_source(source)
        xml_source = _xml_source(source)
        inputs[side] = {
            "format": "sqlite"
            if sqlite_source
            else "xlsx"
            if xlsx_source
            else "json"
            if json_source
            else "xml"
            if xml_source
            else _file_format(paths[0]),
            "bytes": sizes[side],
            "columns": len(columns),
            "partitions": len(paths),
            "schema": schema,
        }
        if sqlite_source:
            inputs[side]["table"] = sqlite_source[1]
        elif xlsx_source:
            inputs[side].update(_xlsx_resolved_metadata(source))
        elif json_source:
            inputs[side]["pointer"] = json_source[1]
        elif xml_source:
            inputs[side].update(record_path=xml_source[1], namespaces=dict(sorted(xml_source[2].items())))
        if inputs[side]["format"] in {"csv", "tsv"}:
            inputs[side]["delimiter"] = delimiter
            inputs[side]["null_tokens"] = recipe["null_tokens"][side]
        if any(path in decoded_sizes for path in paths):
            inputs[side].update(
                compression="zip" if xlsx_source else "gzip",
                decoded_bytes=sum(decoded_sizes.get(path, 0) for path in paths),
            )
        if validate_records and not schema["missing"] and not schema["unexpected"]:
            inputs[side]["records"] = (
                _aggregate_record_diagnostics(source, recipe, max_rows, max_groups, side, decoded_sizes)
                if recipe["comparison_mode"] == "aggregate"
                else _multiset_record_diagnostics(source, recipe, max_rows, max_distinct_rows, side, decoded_sizes)
                if recipe["comparison_mode"] == "multiset"
                else _record_diagnostics(source, recipe, max_rows, side, decoded_sizes)
            )
            if inputs[side]["records"]["status"] == "invalid":
                status = "invalid"
    if validate_records and any(_source_digest(source) != digest for source, digest in before.items()):
        raise ParisonError("an input changed while it was being validated")
    uses_xlsx = any(_xlsx_source(source) for source in (baseline, candidate))
    uses_xlsx_table = any(
        (selection := _xlsx_source(source)) is not None and selection[1] == "table" for source in (baseline, candidate)
    )
    uses_xlsx_defined_name = any(
        (selection := _xlsx_source(source)) is not None and selection[1] == "defined-name"
        for source in (baseline, candidate)
    )
    uses_json = any(_json_source(source) for source in (baseline, candidate))
    uses_xml = any(_xml_source(source) for source in (baseline, candidate))
    result = {
        "schema_version": (
            23
            if uses_xlsx_defined_name and recipe["recipe_version"] == 5
            else 24
            if uses_xlsx_defined_name and recipe["recipe_version"] == 6
            else 22
            if uses_xlsx_defined_name and recipe["recipe_version"] == 4
            else 17
            if recipe["recipe_version"] == 5
            else 18
            if recipe["recipe_version"] == 6
            else 16
            if recipe["recipe_version"] == 4
            else 20
            if uses_xlsx_defined_name and recipe["comparison_mode"] == "aggregate"
            else 21
            if uses_xlsx_defined_name and recipe["comparison_mode"] == "multiset"
            else 19
            if uses_xlsx_defined_name
            else 14
            if uses_xlsx_table and recipe["comparison_mode"] == "aggregate"
            else 15
            if uses_xlsx_table and recipe["comparison_mode"] == "multiset"
            else 13
            if uses_xlsx_table
            else 11
            if uses_xml and recipe["comparison_mode"] == "aggregate"
            else 12
            if uses_xml and recipe["comparison_mode"] == "multiset"
            else 10
            if uses_xml
            else 8
            if uses_json and recipe["comparison_mode"] == "aggregate"
            else 9
            if uses_json and recipe["comparison_mode"] == "multiset"
            else 7
            if uses_json
            else 5
            if uses_xlsx and recipe["comparison_mode"] == "aggregate"
            else 6
            if uses_xlsx and recipe["comparison_mode"] == "multiset"
            else 4
            if uses_xlsx
            else 2
            if recipe["comparison_mode"] == "aggregate"
            else 3
            if recipe["comparison_mode"] == "multiset"
            else 1
        ),
        "status": status,
        "comparison_mode": recipe["comparison_mode"],
        "canonical_columns": len(recipe["columns"]),
        "policy_sha256": policy_sha256,
        "inputs": inputs,
    }
    if "field_projection" in recipe:
        result["field_projection"] = recipe["field_projection"]
    if recipe["comparison_mode"] == "aggregate":
        result.update(group_by=recipe["group_by"], measures=list(recipe["measures"]), max_groups=max_groups)
    elif recipe["comparison_mode"] == "multiset":
        result["max_distinct_rows"] = max_distinct_rows
    else:
        result["keys"] = recipe["keys"]
    return result


def _suggest_type(values: list[Any]) -> dict[str, Any]:
    present = [str(value) for value in values if value not in (None, "")]
    if not present:
        return {"type": "string", "comparison": "exact"}
    lowered = {value.lower() for value in present}
    if lowered <= {"true", "false"}:
        return {"type": "boolean", "comparison": "exact"}
    if all(
        value.lstrip("-").isdigit() and not (value.lstrip("-").startswith("0") and len(value.lstrip("-")) > 1)
        for value in present
    ):
        return {"type": "integer", "comparison": "exact"}
    try:
        decimals = [Decimal(value) for value in present]
        if all(value.is_finite() for value in decimals):
            scale = max(max(-value.as_tuple().exponent, 0) for value in decimals)
            return {"type": "decimal", "comparison": "exact", "scale": scale}
    except InvalidOperation:
        pass
    try:
        parsed = [datetime.fromisoformat(value.replace("Z", "+00:00")) for value in present]
        if all(value.tzinfo is not None for value in parsed):
            return {"type": "timestamp", "comparison": "exact", "timezone": "require-aware"}
    except ValueError:
        pass
    try:
        if all(date.fromisoformat(value) for value in present):
            return {"type": "date", "comparison": "exact"}
    except ValueError:
        pass
    return {"type": "string", "comparison": "exact"}


def _suggest_keys(columns: list[str], baseline: list[dict[str, Any]], candidate: list[dict[str, Any]]) -> list[str]:
    if not baseline or not candidate:
        return [columns[0]]

    def distinct_counts(names: list[str]) -> list[int]:
        counts = []
        for rows in (baseline, candidate):
            values = [tuple(row[name] for name in names) for row in rows]
            if any(value in (None, "") for key in values for value in key):
                return [-1, -1]
            counts.append(len(set(values)))
        return counts

    identifiers = [
        name
        for name in columns
        if any(token in name.lower().replace("_", " ").split() for token in ("id", "key", "code"))
    ]
    chosen: list[str] = []
    for pool in (identifiers, [name for name in columns if name not in identifiers]):
        remaining = list(pool)
        while remaining:
            best = max(remaining, key=lambda name: min(distinct_counts([*chosen, name])))
            chosen.append(best)
            remaining.remove(best)
            if distinct_counts(chosen) == [len(baseline), len(candidate)]:
                return chosen
    return [columns[0]]


def draft_recipe(
    baseline: str | Path,
    candidate: str | Path,
    max_input_bytes: int = 1_000_000_000,
    max_decoded_bytes: int = 1_000_000_000,
    aggregate: bool = False,
    multiset: bool = False,
) -> dict[str, Any]:
    if aggregate and multiset:
        raise ParisonError("aggregate and multiset drafting are mutually exclusive")
    if max_input_bytes <= 0:
        raise ParisonError("max_input_bytes must be positive")
    for source in (baseline, candidate):
        path = _source_path(source)
        if path.is_symlink():
            raise ParisonError(f"input must not be a symlink: {path}")
        if not path.is_file() and not path.is_dir():
            raise ParisonError(f"input is not a regular file or partition directory: {path}")
        _source_paths(source)
    input_bytes = _source_bytes(baseline) + _source_bytes(candidate)
    if input_bytes > max_input_bytes:
        raise ParisonError(f"combined input size {input_bytes} exceeds limit {max_input_bytes} bytes")
    decoded_sizes = _decoded_sizes((baseline, candidate), max_decoded_bytes)
    before = {str(source): _source_digest(source) for source in (baseline, candidate)}
    delimiters = {
        side: "\t" if _file_format(_source_paths(source)[0]) == "tsv" else ","
        for side, source in (("baseline", baseline), ("candidate", candidate))
    }
    baseline_columns = _input_columns(baseline, decoded_sizes, delimiters["baseline"])
    candidate_columns = _input_columns(candidate, decoded_sizes, delimiters["candidate"])
    if any(_source_digest(source) != digest for source, digest in before.items()):
        raise ParisonError("an input changed while it was being inspected")
    candidate_names = set(candidate_columns)
    shared = [name for name in baseline_columns if name in candidate_names]
    if not shared:
        raise ParisonError("inputs have no shared columns")
    shared_names = set(shared)
    excluded = [name for name in baseline_columns + candidate_columns if name not in shared_names]
    inspection_recipe = {
        "columns": {name: {"type": "string", "comparison": "exact"} for name in shared},
        "excluded_columns": {name: "inspection" for name in excluded},
        "delimiters": delimiters,
        "null_tokens": {"baseline": [], "candidate": []},
    }
    uses_structured = any(
        _xlsx_source(source) or _json_source(source) or _xml_source(source) for source in (baseline, candidate)
    )
    if uses_structured:
        inspection_recipe["input_projection"] = "semantic-scalar-v1"
    if uses_structured:
        baseline_rows = list(_iter_input_rows(baseline, inspection_recipe, 5_000_000, "baseline", decoded_sizes))
        candidate_rows = list(_iter_input_rows(candidate, inspection_recipe, 5_000_000, "candidate", decoded_sizes))
    else:
        baseline_rows, _ = _read(baseline, inspection_recipe, 5_000_000, "baseline", decoded_sizes)
        candidate_rows, _ = _read(candidate, inspection_recipe, 5_000_000, "candidate", decoded_sizes)
    if any(_source_digest(source) != digest for source, digest in before.items()):
        raise ParisonError("an input changed while it was being inspected")
    cutoff = (
        datetime.fromtimestamp(
            max(path.stat().st_mtime for source in (baseline, candidate) for path in _source_paths(source)), UTC
        )
        .isoformat()
        .replace("+00:00", "Z")
    )
    columns = {name: _suggest_type([row[name] for row in baseline_rows + candidate_rows]) for name in shared}
    common = {
        "scope": {
            "snapshot": f"{_source_path(baseline).stem} vs {_source_path(candidate).stem}",
            "cutoff": cutoff,
            "filters": [],
            "completeness": "full",
            "expected_empty": not baseline_rows and not candidate_rows,
        },
        "delimiters": delimiters,
        "null_tokens": {"baseline": [], "candidate": []},
        "columns": columns,
        "column_mappings": {},
        "excluded_columns": {
            name: "present only in baseline" if name in baseline_columns else "present only in candidate"
            for name in excluded
        },
        "output": {"sensitivity": "summary"},
    }
    if uses_structured:
        common["input_projection"] = "semantic-scalar-v1"
    if aggregate:
        rows = baseline_rows + candidate_rows
        group_by = []
        for name in shared:
            values = [row[name] for row in rows]
            if values and all(value not in (None, "") for value in values) and 1 < len(set(values)) < len(values):
                group_by = [name]
                break
        measures = {"rows": {"operator": "count"}}
        for name, policy in columns.items():
            identifier = any(token in name.lower().replace("_", " ").split() for token in ("id", "key", "code"))
            if policy["type"] in {"integer", "decimal"} and name not in group_by and not identifier:
                measures[f"sum_{name}"] = {"operator": "sum", "column": name, "nulls": "reject"}
        return {
            "recipe_version": 2,
            "comparison_mode": "aggregate",
            "group_by": group_by,
            "measures": measures,
            **common,
        }
    if multiset:
        return {"recipe_version": 3, "comparison_mode": "multiset", **common}
    return {
        "recipe_version": 1,
        "comparison_mode": "keyed",
        "keys": _suggest_keys(shared, baseline_rows, candidate_rows),
        "identity": {"null_keys": "reject", "duplicates": "reject"},
        "nulls_equal": True,
        **common,
    }


def _iter_input_rows(
    path: str | Path, recipe: dict[str, Any], max_rows: int, side: str, decoded_sizes: dict[Path, int] | None = None
):
    decoded_sizes = decoded_sizes or {}
    source_path = _source_path(path)
    field_projection = recipe.get("field_projection", {}).get(side)
    _validate_field_projection_source(path, recipe, side)
    if source_path.is_dir():
        row_count = 0
        for partition in _source_paths(path):
            for row in _iter_input_rows(partition, recipe, max_rows, side, decoded_sizes):
                if row_count >= max_rows:
                    raise ParisonError(f"row count in {source_path} exceeds limit {max_rows}")
                row_count += 1
                yield row
        return
    if _sqlite_source(path):
        stream = _sqlite_rows(path, max_rows)
        headers, _ = next(stream)
        _validate_headers(Path(str(path)), headers, recipe, side)
        for _, row in stream:
            yield row
        return
    if _xlsx_source(path):
        if recipe.get("input_projection") != "semantic-scalar-v1":
            raise ParisonError("XLSX inputs require input_projection 'semantic-scalar-v1'")
        stream = _xlsx_rows(path, max_rows)
        headers, _ = next(stream)
        _validate_headers(_source_path(path), headers, recipe, side)
        for _, row in stream:
            yield row
        return
    if _json_source(path):
        if recipe.get("input_projection") != "semantic-scalar-v1":
            raise ParisonError("JSON document inputs require input_projection 'semantic-scalar-v1'")
        selectors = (
            field_projection["fields"] if field_projection and field_projection["kind"] == "json-pointer-v1" else None
        )
        stream = _json_rows(path, max_rows, selectors)
        headers, _ = next(stream)
        _validate_headers(_source_path(path), headers, recipe, side)
        for _, row in stream:
            yield row
        return
    if _xml_source(path):
        if recipe.get("input_projection") != "semantic-scalar-v1":
            raise ParisonError("XML inputs require input_projection 'semantic-scalar-v1'")
        stream = _xml_rows(path, max_rows, field_projection)
        headers, _ = next(stream)
        _validate_headers(_source_path(path), headers, recipe, side)
        for _, row in stream:
            yield row
        return
    path = Path(path)
    if path.is_symlink():
        raise ParisonError(f"input must not be a symlink: {path}")
    if not path.is_file():
        raise ParisonError(f"input is not a regular file: {path}")
    format_name = _file_format(path)
    if format_name in {"csv", "tsv"}:
        try:
            handle = _open_text(path, decoded_sizes, newline="")
        except (OSError, EOFError, gzip.BadGzipFile) as exc:
            raise ParisonError(f"cannot read {path}: {exc}") from exc
        with handle:
            try:
                reader = csv.DictReader(handle, delimiter=recipe.get("delimiters", {}).get(side, ","), strict=True)
                _validate_headers(path, reader.fieldnames or [], recipe, side)
                null_tokens = set(recipe.get("null_tokens", {}).get(side, []))
                for row_count, raw in enumerate(reader):
                    if row_count >= max_rows:
                        raise ParisonError(f"row count in {path} exceeds limit {max_rows}")
                    if None in raw or any(value is None for value in raw.values()):
                        raise ParisonError(f"ragged CSV row {reader.line_num} in {path}")
                    yield {name: None if value in null_tokens else value for name, value in raw.items()}
            except (csv.Error, UnicodeDecodeError) as exc:
                raise ParisonError(f"cannot parse {path}: {exc}") from exc
    elif format_name == "parquet":
        try:
            import polars as pl
        except ImportError as exc:
            raise ParisonError("Parquet support requires: pip install 'parison[parquet]'") from exc
        try:
            scan = pl.scan_parquet(path)
            _validate_parquet_projection(scan.collect_schema(), recipe, side)
            row_count = scan.select(pl.len()).collect().item()
            if row_count > max_rows:
                raise ParisonError(f"row count in {path} exceeds limit {max_rows}")
            frame = pl.read_parquet(path)
        except ParisonError:
            raise
        except Exception as exc:
            raise ParisonError(f"cannot read {path}: {exc}") from exc
        _validate_headers(path, frame.columns, recipe, side)
        yield from frame.iter_rows(named=True)
    elif format_name == "jsonl":
        rows = _iter_jsonl(path, max_rows, decoded_sizes)
        try:
            first = next(rows)
        except StopIteration:
            return
        _validate_headers(path, list(first), recipe, side)
        yield first
        yield from rows
    else:
        raise ParisonError(
            f"unsupported input format for {path}; use .csv, .tsv, .jsonl, .parquet, a JSON, XML, or XLSX locator, "
            "or gzip-compressed text"
        )
